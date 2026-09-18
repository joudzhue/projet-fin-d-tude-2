import re
import sys
import unicodedata
import zipfile
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / ".codex_deps"))

import fitz
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph


REPORT = ROOT / "rapport" / "RAPPORT_PFE_GROD_FINAL_EMSI.docx"
PDF = ROOT / "rapport" / "RAPPORT_PFE_GROD_FINAL_EMSI.pdf"
EMSI_LOGO = ROOT / "rapport" / "emsi-logo.png"
GROD_LOGO = ROOT / "grod-platform-frontend" / "public" / "grod-logo.png"
GREEN = "006B3C"
DARK_GREEN = "003C2A"
COPPER = "B75A22"
GRAY = "4B5563"


def set_font(run, size=12, bold=None, italic=None, color="111111", name="Times New Roman"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def remove_table_borders(table):
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)


def add_bottom_border(paragraph, color=COPPER, size="10"):
    ppr = paragraph._p.get_or_add_pPr()
    pbd = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    pbd.append(bottom)
    ppr.append(pbd)


def extract_emsi_logo():
    with zipfile.ZipFile(REPORT, "r") as source:
        EMSI_LOGO.write_bytes(source.read("word/media/image2.png"))


def find_cover_section_break(doc, dedication):
    node = dedication._p.getprevious()
    while node is not None:
        if node.find("./" + qn("w:pPr") + "/" + qn("w:sectPr")) is not None:
            return node
        node = node.getprevious()
    raise RuntimeError("Saut de section de couverture introuvable")


def move_new_before(doc, anchor, before):
    new = [element for element in doc._element.body.iterchildren() if element not in before]
    for element in new:
        anchor.addprevious(element)


def rebuild_cover(doc):
    dedication = next(p for p in doc.paragraphs if p.text.strip() == "Dédicace")
    section_break = find_cover_section_break(doc, dedication)
    for element in list(doc._element.body.iterchildren()):
        if element is section_break:
            break
        doc._element.body.remove(element)

    before = set(doc._element.body.iterchildren())
    logos = doc.add_table(rows=1, cols=2)
    logos.alignment = WD_TABLE_ALIGNMENT.CENTER
    logos.autofit = False
    logos.columns[0].width = Cm(8.5)
    logos.columns[1].width = Cm(7.0)
    remove_table_borders(logos)
    left = logos.cell(0, 0)
    right = logos.cell(0, 1)
    left.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    right.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run().add_picture(str(EMSI_LOGO), width=Cm(7.2))
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run().add_picture(str(GROD_LOGO), width=Cm(5.2))

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(36)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("RAPPORT DE PROJET DE FIN D’ÉTUDES")
    set_font(r, size=22, bold=True, color=DARK_GREEN)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("5ᵉ année — Ingénierie Informatique et Réseaux")
    set_font(r, size=14, bold=True, color=GRAY)
    add_bottom_border(p)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("SOUS LE THÈME")
    set_font(r, size=11, bold=True, color=COPPER)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("CONCEPTION ET DÉVELOPPEMENT D’UNE PLATEFORME WEB B2B POUR G-ROD")
    set_font(r, size=18, bold=True, color=DARK_GREEN)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Période de stage : du 01 mars 2026 au 31 août 2026")
    set_font(r, size=12, bold=True, color=GRAY)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Réalisé par : ........................................................................................")
    set_font(r, size=12, bold=True)

    info = doc.add_table(rows=1, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info.autofit = False
    info.columns[0].width = Cm(7.3)
    info.columns[1].width = Cm(8.2)
    remove_table_borders(info)
    p = info.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_before = Pt(8)
    for line, bold in [
        ("Encadrement", True),
        ("Tuteur EMSI : ................................", False),
        ("Tuteur entreprise : ........................", False),
    ]:
        r = p.add_run(line)
        set_font(r, size=11.5, bold=bold, color=DARK_GREEN if bold else "111111")
        r.add_break()
    p = info.cell(0, 1).paragraphs[0]
    p.paragraph_format.space_before = Pt(8)
    for line, bold in [
        ("Organisme d’accueil", True),
        ("Morocco Copper Foundry — G-ROD", False),
        ("Activité : .....................................", False),
        ("Adresse : .....................................", False),
    ]:
        r = p.add_run(line)
        set_font(r, size=11.5, bold=bold, color=DARK_GREEN if bold else "111111")
        r.add_break()

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_bottom_border(p, color=GREEN, size="6")
    r = p.add_run("ANNÉE UNIVERSITAIRE 2025–2026")
    set_font(r, size=11, bold=True, color=GRAY)

    move_new_before(doc, section_break, before)


def normalize(text):
    text = unicodedata.normalize("NFKD", text)
    return "".join(ch.lower() for ch in text if ch.isalnum())


def page_lookup(pdf):
    texts = [normalize(page.get_text()) for page in pdf]

    def locate(label, start=10):
        key = normalize(label)
        candidates = [key, key[:80], key[:55], key[:35]]
        for candidate in candidates:
            if len(candidate) < 12:
                continue
            for index in range(start, len(texts)):
                if candidate in texts[index]:
                    return index
        return None

    return locate


def find_index_field_paragraph(doc, heading_text):
    heading = next(p for p in doc.paragraphs if p.text.strip() == heading_text)
    seen = False
    for paragraph in doc.paragraphs:
        if paragraph._p is heading._p:
            seen = True
            continue
        if not seen:
            continue
        codes = [node.text or "" for node in paragraph._p.iter(qn("w:instrText"))]
        if any("TOC" in code for code in codes):
            return paragraph
        if paragraph.text.strip() and paragraph.style.name.startswith("Heading"):
            break
    # La table des matières du modèle est encapsulée dans un bloc Word (w:sdt),
    # non exposé par python-docx. On retire ce résultat provisoire, on conserve
    # le saut de page et on réinsère un vrai champ TOC visible juste avant lui.
    node = heading._p.getnext()
    page_break = None
    while node is not None:
        nxt = node.getnext()
        if node.tag == qn("w:p") and node.findall(".//" + qn("w:br")):
            page_break = node
            break
        if node.tag == qn("w:p"):
            text = "".join(t.text or "" for t in node.iter(qn("w:t"))).strip()
            if text in {"Table des figures", "Liste des tableaux", "Liste des acronymes"}:
                break
        doc._element.body.remove(node)
        node = nxt
    new_p = OxmlElement("w:p")
    if page_break is not None:
        page_break.addprevious(new_p)
    else:
        heading._p.addnext(new_p)
    return Paragraph(new_p, doc._body)


def start_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    run._r.append(instr)
    run = paragraph.add_run()
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run._r.append(separate)


def end_field(paragraph):
    run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def fill_field_results(paragraph, instruction, entries, font_size=9.5):
    paragraph.clear()
    paragraph.paragraph_format.line_spacing = 1.05
    paragraph.paragraph_format.space_after = Pt(0)
    tabs = paragraph.paragraph_format.tab_stops
    tabs.add_tab_stop(Cm(15.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    start_field(paragraph, instruction)
    for index, (label, page, level) in enumerate(entries):
        indent = "    " * max(0, level - 1)
        run = paragraph.add_run(f"{indent}{label}")
        set_font(run, size=font_size, bold=level == 1, color=DARK_GREEN if level == 1 else "111111")
        run = paragraph.add_run(f"\t{page}")
        set_font(run, size=font_size, bold=level == 1, color=DARK_GREEN if level == 1 else "111111")
        if index < len(entries) - 1:
            run.add_break()
    end_field(paragraph)


def renumber_caption_results(doc):
    counters = {}
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        match = re.match(r"^(Figure|Tableau) (\d+)\.\d+\s+—", text)
        if not match:
            continue
        kind, chapter = match.groups()
        key = (kind, chapter)
        counters[key] = counters.get(key, 0) + 1
        after_separator = False
        for run in paragraph._p:
            for field_char in run.findall(qn("w:fldChar")):
                field_type = field_char.get(qn("w:fldCharType"))
                if field_type == "separate":
                    after_separator = True
                elif field_type == "end":
                    after_separator = False
            if after_separator:
                text_nodes = run.findall(qn("w:t"))
                if text_nodes:
                    text_nodes[0].text = str(counters[key])
                    after_separator = False
                    break


def populate_indexes(doc, pdf):
    locate = page_lookup(pdf)
    main_start = locate("Introduction générale", 0)
    if main_start is None:
        raise RuntimeError("Page de départ du corps principal introuvable")

    def display_page(physical_index):
        return str(physical_index - main_start + 1)

    toc_entries = []
    started = False
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text == "Introduction générale":
            started = True
        if not started:
            continue
        include = (
            text == "Introduction générale"
            or bool(re.match(r"^\d+\.\s", text))
            or bool(re.match(r"^\d+\.\d+(?:\.\d+)?\s", text))
            or text.startswith("Conclusion du chapitre")
            or text in {"Conclusion générale et perspectives", "Références bibliographiques", "Annexes"}
            or text.startswith("Annexe ")
        )
        if not include:
            continue
        physical = locate(text)
        if physical is None:
            continue
        if text == "Introduction générale" or re.match(r"^\d+\.\s", text) or text in {"Conclusion générale et perspectives", "Références bibliographiques", "Annexes"}:
            level = 1
        elif re.match(r"^\d+\.\d+\.\d+\s", text):
            level = 3
        else:
            level = 2
        toc_entries.append((text, display_page(physical), level))

    figure_entries = []
    table_entries = []
    counters = {}
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        figure_match = re.match(r"^Figure (\d+)\.\d+\s+—\s+(.+)$", text)
        table_match = re.match(r"^Tableau (\d+)\.\d+\s+—\s+(.+)$", text)
        if figure_match:
            chapter, title = figure_match.groups()
            key = ("Figure", chapter)
            counters[key] = counters.get(key, 0) + 1
            label = f"Figure {chapter}.{counters[key]} — {title}"
            physical = locate(title)
            if physical is not None:
                figure_entries.append((label, display_page(physical), 1))
        elif table_match:
            chapter, title = table_match.groups()
            key = ("Tableau", chapter)
            counters[key] = counters.get(key, 0) + 1
            label = f"Tableau {chapter}.{counters[key]} — {title}"
            physical = locate(title)
            if physical is not None:
                table_entries.append((label, display_page(physical), 1))

    fill_field_results(
        find_index_field_paragraph(doc, "Table des matières"),
        'TOC \\o "1-3" \\h \\z \\u',
        toc_entries,
        9.2,
    )
    fill_field_results(
        find_index_field_paragraph(doc, "Table des figures"),
        'TOC \\h \\z \\t "Figure Caption,1"',
        figure_entries,
        9.2,
    )
    fill_field_results(
        find_index_field_paragraph(doc, "Liste des tableaux"),
        'TOC \\h \\z \\t "Table Caption,1"',
        table_entries,
        9.2,
    )
    print(f"Index visibles : {len(toc_entries)} entrées, {len(figure_entries)} figures, {len(table_entries)} tableaux.")


def main():
    extract_emsi_logo()
    pdf = fitz.open(PDF)
    doc = Document(REPORT)
    rebuild_cover(doc)
    renumber_caption_results(doc)
    populate_indexes(doc, pdf)
    doc.save(REPORT)
    print(REPORT)


if __name__ == "__main__":
    main()
