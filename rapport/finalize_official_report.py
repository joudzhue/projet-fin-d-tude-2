import copy
import os
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
DEPS = ROOT / "rapport" / ".docx_deps"
sys.path.insert(0, str(DEPS))

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docxcompose.composer import Composer


SOURCE = ROOT / "rapport" / "RAPPORT_PFE_GROD_VERSION_1_CORRIGEE.docx"
COVER = Path(r"C:\Users\hp\Downloads\Page de garde Stage 5IIR.docx")
LOGO = ROOT / "grod-platform-frontend" / "public" / "grod-logo.png"
OUTPUT = ROOT / "rapport" / "RAPPORT_PFE_GROD_EMSI_OFFICIEL.docx"

GREEN = "006B3C"
DARK_GREEN = "003C2A"
COPPER = "B75A22"
MID_GRAY = "64706A"
LIGHT_GRAY = "F2F4F3"
BLACK = "111111"


def set_font(run, name="Times New Roman", size=12, bold=None, italic=None, color=BLACK):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def paragraph_text(paragraph):
    return "".join(node.text or "" for node in paragraph._p.iter(qn("w:t")))


def find_paragraph(doc, exact):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == exact:
            return paragraph
    raise ValueError(f"Paragraphe introuvable : {exact}")


def move_before(anchor, elements):
    for element in elements:
        anchor._p.addprevious(element)


def collect_new_body_elements(doc, before_elements):
    return [element for element in doc._element.body.iterchildren() if element not in before_elements]


def append_and_move(doc, anchor, builder):
    before_elements = set(doc._element.body.iterchildren())
    builder()
    elements = collect_new_body_elements(doc, before_elements)
    move_before(anchor, elements)


def add_page_break(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._r.append(br)
    return p


def set_arabic(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ppr = paragraph._p.get_or_add_pPr()
    bidi = ppr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        ppr.append(bidi)
    bidi.set(qn("w:val"), "1")
    for run in paragraph.runs:
        set_font(run, name="Arial", size=12)
        rpr = run._element.get_or_add_rPr()
        rtl = rpr.find(qn("w:rtl"))
        if rtl is None:
            rtl = OxmlElement("w:rtl")
            rpr.append(rtl)
        rtl.set(qn("w:val"), "1")


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="AFC9BB"):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "dashed")
        node.set(qn("w:sz"), "8")
        node.set(qn("w:color"), color)


def add_field(paragraph, instruction, placeholder):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run = OxmlElement("w:r")
        run.append(node)
        paragraph._p.append(run)


def set_page_numbering(sectpr, fmt, start):
    for existing in list(sectpr.findall(qn("w:pgNumType"))):
        sectpr.remove(existing)
    page_num = OxmlElement("w:pgNumType")
    sectpr.append(page_num)
    page_num.set(qn("w:fmt"), fmt)
    page_num.set(qn("w:start"), str(start))
    title_page = sectpr.find(qn("w:titlePg"))
    if title_page is not None:
        sectpr.remove(title_page)


def set_page_geometry(sectpr):
    pgmar = sectpr.find(qn("w:pgMar"))
    if pgmar is None:
        pgmar = OxmlElement("w:pgMar")
        sectpr.append(pgmar)
    pgmar.set(qn("w:top"), "1417")
    pgmar.set(qn("w:bottom"), "1417")
    pgmar.set(qn("w:left"), "1701")
    pgmar.set(qn("w:right"), "1417")


def remove_custom_cover(doc):
    rem = find_paragraph(doc, "Remerciements")
    for element in list(doc._element.body.iterchildren()):
        if element is rem._p:
            break
        doc._element.body.remove(element)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)

    tokens = {
        "Title": (18, DARK_GREEN),
        "Heading 1": (16, DARK_GREEN),
        "Heading 2": (14, GREEN),
        "Heading 3": (12, COPPER),
    }
    for name, (size, color) in tokens.items():
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)

    caption = doc.styles["Caption"]
    caption.font.name = "Times New Roman"
    caption.font.size = Pt(10)
    caption.font.italic = False
    caption.paragraph_format.line_spacing = 1.0
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for style_name in ("Figure Caption", "Table Caption"):
        if style_name not in doc.styles:
            style = doc.styles.add_style(style_name, 1)
        else:
            style = doc.styles[style_name]
        style.base_style = caption
        style.font.name = "Times New Roman"
        style.font.size = Pt(10)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style.paragraph_format.line_spacing = 1.0


def normalize_headings_and_captions(doc):
    chapter_titles = {
        "Contexte général et environnement du projet": "1. Contexte général et environnement du projet",
        "Analyse de l'existant et cadrage du projet": "2. Analyse de l'existant et cadrage du projet",
        "Étude, analyse et conception de la solution": "3. Étude, analyse et conception de la solution",
        "Réalisation et mise en œuvre": "4. Réalisation et mise en œuvre",
    }
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text in chapter_titles:
            paragraph.text = chapter_titles[text]
            paragraph.style = doc.styles["Heading 1"]
            for run in paragraph.runs:
                set_font(run, size=16, bold=True, color=DARK_GREEN)
        elif re.match(r"^\d+\.\d+\.\d+\s", text):
            paragraph.style = doc.styles["Heading 3"]
        elif re.match(r"^\d+\.\d+\s", text):
            paragraph.style = doc.styles["Heading 2"]
        elif text.startswith("Figure "):
            paragraph.style = doc.styles["Figure Caption"]
        elif text.startswith("Tableau "):
            paragraph.style = doc.styles["Table Caption"]


def insert_dedication(doc):
    anchor = find_paragraph(doc, "Remerciements")

    def builder():
        p = doc.add_paragraph("Dédicace", style="Heading 1")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(48)
        p.paragraph_format.space_after = Pt(18)
        r = p.add_run("À ma famille, à mes enseignants et à toutes les personnes qui m’ont soutenu tout au long de mon parcours.")
        set_font(r, size=12, italic=True, color=MID_GRAY)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("[Vous pouvez personnaliser cette dédicace avant le dépôt final.]")
        set_font(r, size=10, italic=True, color=COPPER)
        add_page_break(doc)

    append_and_move(doc, anchor, builder)


def insert_arabic_summary(doc):
    anchor = find_paragraph(doc, "Table des matières")

    def builder():
        p = doc.add_paragraph("ملخص", style="Heading 1")
        set_arabic(p)
        paragraphs = [
            "يهدف مشروع نهاية الدراسة هذا إلى تصميم وتطوير منصة ويب مهنية لفائدة شركة Morocco Copper Foundry تحت علامة G-ROD. جاءت المنصة استجابة للحاجة إلى تقديم منتجات النحاس بصورة حديثة، وتسهيل طلبات عروض الأسعار والوثائق التقنية، ثم تجميع معالجتها داخل فضاء إدارة آمن ومنظم.",
            "تتكون المنصة من موقع عام متجاوب يضم كتالوجاً ديناميكياً وصفحات تفصيلية للمنتجات ومعاينات ثلاثية الأبعاد وفضاءً للموارد التقنية، إضافة إلى نماذج للتواصل وطلب عرض السعر. كما يتيح فضاء الإدارة متابعة الطلبات ضمن مسار تجاري، وإدارة المنتجات والعملاء والوثائق والإشعارات. تم تطوير الواجهة الخلفية باستعمال Java 21 وSpring Boot، والواجهة الأمامية باستعمال React وVite، مع حفظ البيانات في قاعدة MySQL.",
            "تعتمد الحماية على Spring Security والمصادقة بواسطة JWT وتشفير كلمات المرور باستخدام BCrypt وتحديد عدد الطلبات، مع دعم اختياري لمفاتيح المرور WebAuthn. وتم التحقق من جودة الحل عن طريق اختبارات وحدات وتكامل واختبارات شاملة باستعمال Maven وH2 وPlaywright. أظهرت النتائج أن النسخة المرشحة للإصدار عملية ومتناسقة وقابلة للتحضير للنشر خلف Nginx، مع إمكانية تطويرها لاحقاً عبر تحسين تقسيم الواجهة وأتمتة النشر.",
        ]
        for text in paragraphs:
            p = doc.add_paragraph(text)
            set_arabic(p)
        p = doc.add_paragraph("الكلمات المفتاحية: منصة مهنية، النحاس، React، Spring Boot، الأمن.")
        set_arabic(p)
        add_page_break(doc)

    append_and_move(doc, anchor, builder)


def replace_keywords(doc):
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("Mots-clés :"):
            paragraph.clear()
            r = paragraph.add_run("Mots-clés : ")
            set_font(r, bold=True, color=DARK_GREEN)
            r = paragraph.add_run("plateforme B2B, cuivre, React, Spring Boot, sécurité.")
            set_font(r)
        elif text.startswith("Keywords:"):
            paragraph.clear()
            r = paragraph.add_run("Keywords: ")
            set_font(r, bold=True, color=DARK_GREEN)
            r = paragraph.add_run("B2B platform, copper, React, Spring Boot, security.")
            set_font(r)


def replace_list_with_field(doc, heading_text, instruction, placeholder):
    heading = find_paragraph(doc, heading_text)
    current = heading._p.getnext()
    while current is not None:
        next_node = current.getnext()
        has_break = bool(current.xpath('.//w:br[@w:type="page"]'))
        if has_break:
            break
        doc._element.body.remove(current)
        current = next_node
    before_elements = set(doc._element.body.iterchildren())
    p = doc.add_paragraph()
    add_field(p, instruction, placeholder)
    elements = collect_new_body_elements(doc, before_elements)
    move_before(type("Anchor", (), {"_p": current})(), elements)


def insert_main_section_break(doc):
    intro = find_paragraph(doc, "Introduction générale")
    body_sectpr = doc._element.body.sectPr
    prelim_sectpr = copy.deepcopy(body_sectpr)
    set_page_geometry(prelim_sectpr)
    set_page_numbering(prelim_sectpr, "lowerRoman", 1)
    p = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    ppr.append(prelim_sectpr)
    p.append(ppr)
    intro._p.addprevious(p)
    set_page_geometry(body_sectpr)
    set_page_numbering(body_sectpr, "decimal", 1)


def add_capture_placeholder(doc, anchor, caption, description):
    def builder():
        add_page_break(doc)
        p = doc.add_paragraph(
            f"La {caption.split(' — ')[0].lower()} présente {description.lower()}. "
            "Elle sera remplacée par la capture ou le diagramme final avant le dépôt."
        )
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.columns[0].width = Cm(15.2)
        row = table.rows[0]
        row.height = Cm(11.2)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        cell = row.cells[0]
        cell.width = Cm(15.2)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, LIGHT_GRAY)
        set_cell_borders(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run("ZONE RÉSERVÉE À LA CAPTURE / AU DIAGRAMME")
        set_font(r, size=12, bold=True, color=GREEN)
        p = cell.add_paragraph("Insérez ici une image nette, recadrée et lisible.")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.runs[0]
        set_font(r, size=10, italic=True, color=MID_GRAY)
        cap = doc.add_paragraph(caption, style="Figure Caption")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    append_and_move(doc, anchor, builder)


def insert_capture_pages(doc):
    groups = {
        "Conclusion du chapitre 2": [
            ("Figure 2.1 — Planification prévisionnelle du projet", "le diagramme de Gantt et les principales phases du projet"),
        ],
        "Conclusion du chapitre 3": [
            ("Figure 3.1 — Architecture globale de la plateforme G-ROD", "les couches frontend, API, base de données et stockage documentaire"),
            ("Figure 3.2 — Diagramme de cas d’utilisation du site public", "les interactions du visiteur et du client avec le site public"),
            ("Figure 3.3 — Diagramme de cas d’utilisation de l’espace Admin", "les fonctions accessibles à l’administrateur"),
            ("Figure 3.4 — Diagramme de séquence d’une demande de devis", "le parcours complet de création d’une demande de devis"),
            ("Figure 3.5 — Diagramme de séquence du traitement administratif", "la consultation et la mise à jour du statut d’une demande"),
            ("Figure 3.6 — Modèle conceptuel des données", "les entités métier et leurs principales relations"),
        ],
        "Conclusion du chapitre 4": [
            ("Figure 4.3 — Page d’accueil publique", "la page d’accueil de la plateforme dans sa version finale"),
            ("Figure 4.4 — Catalogue en mode cartes et en mode liste", "les deux modes d’affichage du catalogue"),
            ("Figure 4.5 — Fiche produit et aperçu 3D", "la fiche détaillée d’un produit et son visualiseur tridimensionnel"),
            ("Figure 4.6 — Formulaire de demande de devis", "les champs du formulaire public et ses contrôles"),
            ("Figure 4.7 — Ressources et demande de document", "la consultation et la demande d’une ressource technique"),
            ("Figure 4.8 — Affichage responsive du site public", "le rendu mobile ou tablette des principales pages publiques"),
            ("Figure 4.9 — Dashboard commercial Admin", "les indicateurs et raccourcis du tableau de bord"),
            ("Figure 4.10 — Pipeline commercial des demandes", "la répartition paginée des demandes selon leur statut"),
            ("Figure 4.11 — Gestion des demandes", "la liste, les filtres et les actions de traitement des demandes"),
            ("Figure 4.12 — Gestion des clients", "le tableau des clients et l’accès à leur historique"),
            ("Figure 4.13 — Gestion des produits", "le catalogue administratif et le formulaire produit"),
            ("Figure 4.14 — Gestion des demandes de documents", "le suivi des demandes de documents techniques"),
            ("Figure 4.15 — Gestion des ressources techniques", "la liste des ressources et le formulaire d’ajout"),
            ("Figure 4.16 — Centre de notifications", "le panneau et la page des notifications persistantes"),
            ("Figure 4.17 — Mon compte et sécurité", "les onglets Profil, Sécurité et Préférences"),
            ("Figure 4.18 — Sidebar réduite et drawer mobile", "le comportement responsive de la navigation Admin"),
            ("Figure 4.19 — Installation de la PWA", "le parcours d’installation de la plateforme sur un appareil"),
            ("Figure 4.20 — Résultats des tests automatisés", "la synthèse d’exécution des tests backend et frontend"),
        ],
    }
    for anchor_text, figures in groups.items():
        anchor = find_paragraph(doc, anchor_text)
        for caption, description in figures:
            add_capture_placeholder(doc, anchor, caption, description)


def add_chapter_introductions(doc):
    chapter_titles = [
        ("1. Contexte général et environnement du projet", "Ce chapitre présente l’organisme d’accueil, son activité industrielle, la gamme G-ROD et le contexte ayant motivé la réalisation de la plateforme."),
        ("2. Analyse de l'existant et cadrage du projet", "Ce chapitre analyse l’existant, précise les acteurs, formalise les besoins fonctionnels et non fonctionnels, puis expose l’organisation retenue pour conduire le projet."),
        ("3. Étude, analyse et conception de la solution", "Ce chapitre justifie les technologies retenues et décrit l’architecture, les cas d’utilisation, le modèle de données ainsi que les principaux choix de sécurité."),
        ("4. Réalisation et mise en œuvre", "Ce chapitre présente l’environnement de développement, les modules réalisés, les mécanismes de sécurité et de gestion des fichiers, la recette ainsi que la préparation du déploiement."),
    ]
    for title, intro_text in chapter_titles:
        chapter = find_paragraph(doc, title)
        current = chapter._p.getnext()
        while current is not None and current.tag != qn("w:p"):
            current = current.getnext()
        # The existing chapter title page ends with a page break. Insert the
        # introduction immediately before the first numbered subsection.
        target = None
        seen = False
        for paragraph in doc.paragraphs:
            if paragraph._p is chapter._p:
                seen = True
                continue
            if seen and re.match(r"^\d+\.\d+\s", paragraph.text.strip()):
                target = paragraph
                break
        if target is None:
            continue

        def builder(text=intro_text):
            p = doc.add_paragraph("Introduction du chapitre", style="Heading 2")
            p = doc.add_paragraph(text)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        append_and_move(doc, target, builder)


def add_citations_and_references(doc):
    citation_rules = [
        ("Morocco Copper Foundry, abrégée MCF", " [1]"),
        ("La documentation institutionnelle situe la création de MCF", " [1]"),
        ("L'architecture de production documentée place Nginx", " [2]"),
        ("Le backend repose sur Java 21", " [3]"),
        ("Le frontend repose sur React", " [4]"),
        ("MySQL", " [5]"),
    ]
    used_starts = set()
    for paragraph in doc.paragraphs:
        text = paragraph.text
        for needle, suffix in citation_rules:
            if needle in text and needle not in used_starts and not text.rstrip().endswith(suffix.strip()):
                paragraph.add_run(suffix)
                used_starts.add(needle)
                break

    heading = find_paragraph(doc, "Bibliographie et webographie")
    heading.text = "Références bibliographiques"
    heading.style = doc.styles["Heading 1"]
    current = heading._p.getnext()
    while current is not None:
        next_node = current.getnext()
        if current.tag == qn("w:p"):
            text = paragraph_text(type("P", (), {"_p": current})())
            if text.strip() == "Annexes":
                break
        has_break = bool(current.xpath('.//w:br[@w:type="page"]'))
        if has_break:
            doc._element.body.remove(current)
            break
        doc._element.body.remove(current)
        current = next_node

    annexes = find_paragraph(doc, "Annexes")

    def builder():
        refs = [
            "[1] Morocco Copper Foundry, Brochure institutionnelle G-ROD, document d’entreprise intégré au projet, 2026.",
            "[2] Projet G-ROD, README et procédures de déploiement, sauvegarde et migration, version v1.0.0-rc.2, 2026.",
        ]
        for ref in refs:
            p = doc.add_paragraph(ref)
            p.paragraph_format.left_indent = Cm(0.7)
            p.paragraph_format.first_line_indent = Cm(-0.7)
        p = doc.add_paragraph("Webographie", style="Heading 1")
        sites = [
            "[3] Spring, Spring Boot Reference Documentation, https://docs.spring.io/spring-boot/ (consulté le 12/09/2026).",
            "[4] Meta, React Documentation, https://react.dev/ (consultée le 12/09/2026).",
            "[5] Oracle, MySQL 8.0 Reference Manual, https://dev.mysql.com/doc/refman/8.0/en/ (consulté le 12/09/2026).",
            "[6] Vite, Guide officiel, https://vite.dev/guide/ (consulté le 12/09/2026).",
            "[7] Three.js, Documentation officielle, https://threejs.org/docs/ (consultée le 12/09/2026).",
            "[8] Microsoft, Playwright Documentation, https://playwright.dev/docs/intro (consultée le 12/09/2026).",
            "[9] OWASP Foundation, Cheat Sheet Series, https://cheatsheetseries.owasp.org/ (consultée le 12/09/2026).",
        ]
        for site in sites:
            p = doc.add_paragraph(site)
            p.paragraph_format.left_indent = Cm(0.7)
            p.paragraph_format.first_line_indent = Cm(-0.7)
        add_page_break(doc)

    append_and_move(doc, annexes, builder)


def set_document_geometry(doc):
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2.5)
        section.header_distance = Cm(1.25)
        section.footer_distance = Cm(1.25)
        section.different_first_page_header_footer = False


def patch_cover_and_logo(path):
    replacements = {
        "TITRE DU THEME": "CONCEPTION ET DÉVELOPPEMENT D’UNE PLATEFORME WEB B2B POUR G-ROD",
        "Année UNIVERSITAIRE: 2025-2026": "ANNÉE UNIVERSITAIRE : 2025-2026",
    }
    temp_path = path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "word/document.xml":
                xml = data.decode("utf-8")
                for old, new in replacements.items():
                    xml = xml.replace(old, new)
                xml = xml.replace("<w:t>Nom</w:t>", "<w:t>Morocco Copper Foundry — G-ROD</w:t>")
                data = xml.encode("utf-8")
            elif info.filename == "word/media/image1.png" and LOGO.exists():
                data = LOGO.read_bytes()
            target.writestr(info, data)
    temp_path.replace(path)


def configure_header_footer(section, numbered=True):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    header = section.header
    for paragraph in header.paragraphs:
        paragraph.clear()
    if numbered:
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run("RAPPORT DE PROJET DE FIN D’ÉTUDES  |  G-ROD")
        set_font(r, name="Arial", size=8.5, bold=True, color=MID_GRAY)
        ppr = p._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), COPPER)
        borders.append(bottom)
        ppr.append(borders)

    footer = section.footer
    for paragraph in footer.paragraphs:
        paragraph.clear()
    if numbered:
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("G-ROD — Morocco Copper Foundry     |     Page ")
        set_font(r, name="Arial", size=8.5, color=MID_GRAY)
        add_field(p, "PAGE", "1")


def normalize_table_geometry(doc, total_width=8788):
    for table in doc.tables:
        tblpr = table._tbl.tblPr
        tblw = tblpr.find(qn("w:tblW"))
        if tblw is None:
            tblw = OxmlElement("w:tblW")
            tblpr.append(tblw)
        tblw.set(qn("w:type"), "dxa")
        tblw.set(qn("w:w"), str(total_width))

        tblind = tblpr.find(qn("w:tblInd"))
        if tblind is None:
            tblind = OxmlElement("w:tblInd")
            tblpr.append(tblind)
        tblind.set(qn("w:type"), "dxa")
        tblind.set(qn("w:w"), "120")

        layout = tblpr.find(qn("w:tblLayout"))
        if layout is None:
            layout = OxmlElement("w:tblLayout")
            tblpr.append(layout)
        layout.set(qn("w:type"), "fixed")

        grid = table._tbl.tblGrid
        grid_columns = list(grid)
        count = max(1, len(grid_columns))
        original = [int(column.get(qn("w:w"), "1")) for column in grid_columns]
        if not original or sum(original) <= 0:
            original = [1] * count
        widths = [round(value * total_width / sum(original)) for value in original]
        widths[-1] += total_width - sum(widths)
        for column, width in zip(grid_columns, widths):
            column.set(qn("w:w"), str(width))
        for row in table.rows:
            for index, cell in enumerate(row.cells):
                width = widths[min(index, len(widths) - 1)]
                tcpr = cell._tc.get_or_add_tcPr()
                tcw = tcpr.find(qn("w:tcW"))
                if tcw is None:
                    tcw = OxmlElement("w:tcW")
                    tcpr.append(tcw)
                tcw.set(qn("w:type"), "dxa")
                tcw.set(qn("w:w"), str(width))


def finalize_sections(path):
    doc = Document(path)
    dedication = find_paragraph(doc, "Dédicace")
    cover_sectpr = copy.deepcopy(doc._element.body.sectPr)
    for node in list(cover_sectpr.findall(qn("w:pgNumType"))):
        cover_sectpr.remove(node)
    p = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    ppr.append(cover_sectpr)
    p.append(ppr)
    dedication._p.addprevious(p)
    doc.save(path)

    doc = Document(path)
    if len(doc.sections) != 3:
        raise RuntimeError(f"Trois sections attendues (couverture, préliminaires, corps), obtenu : {len(doc.sections)}")

    cover, prelim, main = doc.sections
    cover.start_type = WD_SECTION_START.NEW_PAGE
    configure_header_footer(cover, numbered=False)
    for section in (prelim, main):
        section.start_type = WD_SECTION_START.NEW_PAGE
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2.5)
        section.header_distance = Cm(1.25)
        section.footer_distance = Cm(1.25)
        section.different_first_page_header_footer = False
        configure_header_footer(section, numbered=True)
    set_page_numbering(prelim._sectPr, "lowerRoman", 1)
    set_page_numbering(main._sectPr, "decimal", 1)
    normalize_table_geometry(doc)
    doc.save(path)


def patch_accessibility_metadata(path):
    temp_path = path.with_suffix(".a11y.docx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "word/document.xml":
                from lxml import etree

                root = etree.fromstring(data)
                ns = {
                    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
                    "v": "urn:schemas-microsoft-com:vml",
                }
                for idx, node in enumerate(root.xpath("//wp:docPr", namespaces=ns), 1):
                    if not node.get("descr"):
                        name = node.get("name") or f"Élément graphique {idx}"
                        node.set("descr", f"Élément de mise en page : {name}")
                for idx, node in enumerate(root.xpath("//v:shape", namespaces=ns), 1):
                    if not node.get("alt"):
                        node.set("alt", f"Élément graphique de la page de garde {idx}")
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            target.writestr(info, data)
    temp_path.replace(path)


def set_update_fields(doc):
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if not COVER.exists():
        raise FileNotFoundError(COVER)

    content = Document(SOURCE)
    remove_custom_cover(content)
    configure_styles(content)
    normalize_headings_and_captions(content)
    insert_dedication(content)
    insert_arabic_summary(content)
    replace_keywords(content)
    replace_list_with_field(
        content,
        "Table des figures",
        'TOC \\h \\z \\t "Figure Caption,1"',
        "Mettre à jour la table des figures dans Word avec Ctrl+A puis F9.",
    )
    replace_list_with_field(
        content,
        "Liste des tableaux",
        'TOC \\h \\z \\t "Table Caption,1"',
        "Mettre à jour la liste des tableaux dans Word avec Ctrl+A puis F9.",
    )
    add_chapter_introductions(content)
    insert_capture_pages(content)
    add_citations_and_references(content)
    insert_main_section_break(content)
    set_document_geometry(content)
    set_update_fields(content)
    content.core_properties.title = "Rapport PFE — Plateforme G-ROD"
    content.core_properties.subject = "Conception et développement d’une plateforme web B2B pour Morocco Copper Foundry"
    content.core_properties.author = "[NOM ET PRÉNOM DE L’ÉTUDIANT(E)]"
    content.core_properties.comments = "Version officielle structurée selon le guide de rédaction EMSI."

    with tempfile.TemporaryDirectory(prefix="grod_report_") as tmp:
        content_path = Path(tmp) / "contenu.docx"
        content.save(content_path)
        cover_doc = Document(COVER)
        composer = Composer(cover_doc)
        composer.append(Document(content_path))
        OUTPUT.parent.mkdir(parents=True, exist_ok=True)
        composer.save(OUTPUT)

    patch_cover_and_logo(OUTPUT)
    finalize_sections(OUTPUT)
    patch_accessibility_metadata(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
