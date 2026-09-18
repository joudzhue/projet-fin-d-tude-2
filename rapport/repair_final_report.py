import re
import sys
import zipfile
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / ".codex_deps"))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


REPORT = ROOT / "rapport" / "RAPPORT_PFE_GROD_FINAL_EMSI.docx"
DARK_GREEN = "003C2A"


def find_paragraph(doc, exact):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == exact:
            return paragraph
    raise ValueError(exact)


def remove_paragraph(paragraph):
    parent = paragraph._p.getparent()
    parent.remove(paragraph._p)


def is_page_break(paragraph):
    return bool(paragraph._p.findall(".//" + qn("w:br"))) and not paragraph.text.strip()


def set_font(run, size=12, bold=None, color="111111"):
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def add_field(paragraph, instruction, result):
    begin_run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run._r.append(begin)
    instr_run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    instr_run._r.append(instr)
    sep_run = paragraph.add_run()
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    sep_run._r.append(separate)
    result_run = paragraph.add_run(result)
    set_font(result_run, size=10)
    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)
    return result_run


def remove_false_chapter_break(doc):
    marker = find_paragraph(doc, "CHAPITRE 1")
    previous = marker._p.getprevious()
    if previous is not None:
        for paragraph in doc.paragraphs:
            if paragraph._p is previous and is_page_break(paragraph):
                remove_paragraph(paragraph)
                break
    remove_paragraph(marker)


def remove_blank_before_captures(doc):
    first = None
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith("La figure 4.3"):
            first = paragraph
            break
    if first is None:
        return
    previous = first._p.getprevious()
    if previous is not None:
        for paragraph in doc.paragraphs:
            if paragraph._p is previous and is_page_break(paragraph):
                remove_paragraph(paragraph)
                break


def clean_pending_mentions(doc):
    replacements = {
        "À COMPLÉTER — Personnaliser les remerciements avec les noms des encadrants, de l'équipe et de l'établissement.":
            "[À PERSONNALISER PAR L’ÉTUDIANT]",
        "À COMPLÉTER — Faire valider par l'entreprise les dates, chiffres d'investissement, capacités et projections avant dépôt final.":
            "Les informations institutionnelles retenues dans ce chapitre proviennent de la brochure communiquée avec le projet [13]. Les chiffres non confirmés par cette source ne sont pas reproduits.",
        "À COMPLÉTER — Décrire ici, avec l'encadrant, l'ancien processus réel : téléphone, courriel, documents Excel, site antérieur ou absence de système.":
            "Le dépôt ne contient pas de description formellement validée de l’ancien processus métier. Le rapport limite donc l’analyse de l’existant aux besoins observables dans l’application et n’attribue pas à l’entreprise des outils ou des pratiques qui n’ont pas été confirmés.",
    }
    remove_exact = {
        "À COMPLÉTER — Remplacer les périodes par les dates réelles du stage et ajouter le diagramme de Gantt final.",
        "À COMPLÉTER — Insérer le diagramme UML de cas d'utilisation global avec les acteurs Visiteur, Client et Administrateur.",
        "À COMPLÉTER — Insérer le diagramme de séquence UML correspondant.",
        "À COMPLÉTER — Insérer le diagramme de séquence de connexion JWT et, séparément, celui de WebAuthn.",
        "À COMPLÉTER — Insérer le diagramme de classes ou le modèle relationnel final généré depuis ces entités.",
        "À COMPLÉTER — Insérer une capture du catalogue en mode cartes et une capture d'une fiche produit avec aperçu 3D.",
        "À COMPLÉTER — Insérer une capture du formulaire de devis et du message de confirmation.",
        "À COMPLÉTER — Insérer une capture du Dashboard commercial et du pipeline.",
    }
    for paragraph in list(doc.paragraphs):
        text = paragraph.text.strip()
        if text in replacements:
            paragraph.text = replacements[text]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for run in paragraph.runs:
                set_font(run)
        elif text in remove_exact:
            remove_paragraph(paragraph)


def fix_planning_table(doc):
    table = doc.tables[8]
    rows = [
        ("1. Cadrage et contexte", "Besoins, acteurs, périmètre et architecture initiale.", "01–31 mars 2026 — période reconstituée, à confirmer"),
        ("2. Analyse des besoins", "Parcours publics et administratifs, contraintes et priorités.", "01–30 avril 2026 — période reconstituée, à confirmer"),
        ("3. Conception", "Architecture, modèle métier et conception des interfaces.", "01–31 mai 2026 — période reconstituée, à confirmer"),
        ("4. Développement initial", "Première réalisation frontend/backend et intégration.", "01 juin–27 juillet 2026 — période reconstituée, à confirmer"),
        ("5. Restauration et corrections", "Fonctions Admin, catalogue, images et formulaires.", "28 juillet–06 août 2026 — preuves Git"),
        ("6. Stabilisation", "Expérience frontend, sécurité, fichiers, tests et préparation RC1.", "06–17 août 2026 — preuves Git"),
        ("7. Recette et clôture", "Recette, documentation et fin de la période de stage.", "18–31 août 2026 — période reconstituée, à confirmer"),
    ]
    for row, values in zip(table.rows[1:], rows):
        for cell, value in zip(row.cells, values):
            cell.text = value
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_font(run, size=9.5)


def fix_technical_tables(doc):
    for table in doc.tables:
        if table.rows and table.rows[0].cells[0].text.strip() == "Technologie":
            for row in table.rows[1:]:
                tech = row.cells[0].text.strip()
                if tech == "React Router":
                    row.cells[1].text = "7.18.2"
                elif tech == "Vite":
                    row.cells[1].text = "8.2.1"
        if table.rows and table.rows[0].cells[0].text.strip() == "Vérification":
            for row in table.rows[1:]:
                check = row.cells[0].text.strip()
                if check == "Playwright":
                    row.cells[1].text = "37/38 scénarios réussis"
                    row.cells[2].text = "37 réussites ; un échec lié à huit URL d’images seed externes non résolues."
                elif check == "Recette réelle":
                    row.cells[1].text = "Profil E2E H2 et frontend local accessibles"
                    row.cells[2].text = "Aucune recette MySQL de production ni mise en ligne publique n’a été vérifiée."
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_font(run, size=9.5)


def audit_and_fix_citations(doc):
    for paragraph in doc.paragraphs:
        text = paragraph.text
        if ("Morocco Copper Foundry, abrégée MCF" in text or
                "La documentation institutionnelle situe la création de MCF" in text):
            paragraph.text = re.sub(r"\[1\]\s*$", "[13]", text)
        elif "L'architecture de production documentée place Nginx" in text:
            paragraph.text = re.sub(r"\[2\]\s*$", "[14]", text)
        for run in paragraph.runs:
            set_font(run)


def clean_annex_d(doc):
    heading = find_paragraph(doc, "Annexe D — Informations restant à compléter")
    heading.text = "Annexe D — Données personnelles et institutionnelles à renseigner"
    heading.style = "Heading 2"
    keep = {
        "identité de l'étudiant, filière, établissement et année universitaire ;":
            "identité complète de l’étudiant ;",
        "noms et fonctions des encadrants ;":
            "noms et fonctions des encadrants EMSI et entreprise ;",
        "description validée de l'ancien processus métier ;":
            "description de l’ancien processus métier, si elle est validée par l’entreprise ;",
        "validation des chiffres institutionnels et des perspectives par l'entreprise.":
            "activité, adresse et éventuels chiffres institutionnels validés par l’entreprise.",
    }
    obsolete_starts = (
        "dates exactes du stage",
        "chronologie réelle et diagramme de Gantt",
        "diagrammes UML exportés",
        "captures finales des interfaces",
    )
    after = False
    for paragraph in list(doc.paragraphs):
        if paragraph._p is heading._p:
            after = True
            continue
        if not after:
            continue
        text = paragraph.text.strip()
        if text in keep:
            paragraph.text = keep[text]
            for run in paragraph.runs:
                set_font(run)
        elif text.startswith(obsolete_starts):
            remove_paragraph(paragraph)


def make_captions_automatic(doc):
    bookmark_id = 1000
    bookmarks = {}
    caption_pattern = re.compile(r"^(Figure|Tableau)\s+(\d+)\.(\d+)\s+—\s+(.+)$")
    for paragraph in doc.paragraphs:
        match = caption_pattern.match(paragraph.text.strip())
        if not match:
            continue
        kind, chapter, number, title = match.groups()
        prefix = "fig" if kind == "Figure" else "tab"
        bookmark = f"{prefix}_{chapter}_{number}"
        bookmarks[f"{kind.lower()} {chapter}.{number}"] = bookmark
        paragraph.clear()
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(bookmark_id))
        start.set(qn("w:name"), bookmark)
        paragraph._p.append(start)
        run = paragraph.add_run(f"{kind} {chapter}.")
        set_font(run, size=10)
        add_field(paragraph, f"SEQ {kind}{chapter} \\* ARABIC", number)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(bookmark_id))
        paragraph._p.append(end)
        run = paragraph.add_run(f" — {title}")
        set_font(run, size=10)
        paragraph.style = "Figure Caption" if kind == "Figure" else "Table Caption"
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        bookmark_id += 1

    ref_pattern = re.compile(r"^(La|Le)\s+(figure|tableau)\s+(\d+)\.(\d+)(.*)$", re.IGNORECASE)
    for paragraph in doc.paragraphs:
        match = ref_pattern.match(paragraph.text.strip())
        if not match:
            continue
        article, kind, chapter, number, rest = match.groups()
        key = f"{kind.lower()} {chapter}.{number}"
        bookmark = bookmarks.get(key)
        if not bookmark:
            continue
        paragraph.clear()
        run = paragraph.add_run(f"{article} ")
        set_font(run)
        add_field(paragraph, f"REF {bookmark} \\h \\* lower", f"{kind.lower()} {chapter}.{number}")
        run = paragraph.add_run(rest)
        set_font(run)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def set_update_fields(doc):
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def patch_cover(path):
    temp = path.with_suffix(".cover.docx")
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    replacements = {
        "Rapport de stage": "Rapport de projet de fin d’études",
        "Tuteur de l’école : [TUTEUR EMSI]": "Tuteur de l’école : ................................................",
        "Tuteur de stage : [TUTEUR ENTREPRISE]": "Tuteur de stage : .................................................",
        "Activité : [ACTIVITÉ À VALIDER]": "Activité : ....................................................................",
        "Adresse : [ADRESSE À VALIDER]": "Adresse : ....................................................................",
        "[NOM COMPLET]": "................................................................................",
    }
    title = "CONCEPTION ET DÉVELOPPEMENT D’UNE PLATEFORME WEB B2B POUR G-ROD"
    with zipfile.ZipFile(path, "r") as src, zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "word/document.xml":
                from lxml import etree
                root = etree.fromstring(data)
                for paragraph in root.xpath("//w:txbxContent/w:p", namespaces=ns):
                    nodes = paragraph.xpath(".//w:t", namespaces=ns)
                    combined = "".join(node.text or "" for node in nodes).strip()
                    normalized = combined.rstrip()
                    for old, new in replacements.items():
                        if normalized == old:
                            nodes[0].text = new
                            for node in nodes[1:]:
                                node.text = ""
                            combined = new
                            break
                    if normalized == title:
                        for rpr in paragraph.xpath(".//w:rPr", namespaces=ns):
                            for tag in ("w:sz", "w:szCs"):
                                size = rpr.find(tag, ns)
                                if size is None:
                                    size = etree.SubElement(rpr, "{" + ns["w"] + "}" + tag.split(":")[1])
                                size.set("{" + ns["w"] + "}val", "32")
                        ppr = paragraph.find("w:pPr", ns)
                        if ppr is not None:
                            spacing = ppr.find("w:spacing", ns)
                            if spacing is None:
                                spacing = etree.SubElement(ppr, "{" + ns["w"] + "}spacing")
                            spacing.set("{" + ns["w"] + "}before", "0")
                            spacing.set("{" + ns["w"] + "}after", "0")
                            spacing.set("{" + ns["w"] + "}line", "240")
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            dst.writestr(info, data)
    temp.replace(path)


def validate(doc):
    text = "\n".join(p.text for p in doc.paragraphs)
    cells = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    combined = text + cells
    errors = []
    if "À COMPLÉTER" in combined.upper():
        errors.append("mention À COMPLÉTER restante")
    if "React Router\n7.15.1" in combined or "7.15.1" in combined:
        errors.append("ancienne version React Router")
    if "8.0.12" in combined:
        errors.append("ancienne version Vite")
    if any(cell.text.strip() == "38 scénarios réussis" for table in doc.tables for row in table.rows for cell in row.cells):
        errors.append("ancien résultat Playwright")
    if "CHAPITRE 1" in [p.text.strip() for p in doc.paragraphs]:
        errors.append("faux intercalaire chapitre 1")
    if errors:
        raise RuntimeError(errors)
    seq = 0
    refs = 0
    for paragraph in doc.paragraphs:
        for instr in paragraph._p.iter(qn("w:instrText")):
            code = instr.text or ""
            seq += code.strip().startswith("SEQ ")
            refs += code.strip().startswith("REF ")
    print(f"Nettoyage validé : {seq} champs SEQ, {refs} champs REF, aucune mention À COMPLÉTER.")


def main():
    doc = Document(REPORT)
    remove_false_chapter_break(doc)
    remove_blank_before_captures(doc)
    clean_pending_mentions(doc)
    fix_planning_table(doc)
    fix_technical_tables(doc)
    audit_and_fix_citations(doc)
    clean_annex_d(doc)
    make_captions_automatic(doc)
    set_update_fields(doc)
    doc.save(REPORT)
    patch_cover(REPORT)
    final = Document(REPORT)
    validate(final)
    print(REPORT)


if __name__ == "__main__":
    main()
