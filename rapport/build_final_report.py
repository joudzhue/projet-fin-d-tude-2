import copy
import re
import sys
import zipfile
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / ".codex_deps"))

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph
from PIL import Image


SOURCE = ROOT / "rapport" / "RAPPORT_PFE_GROD_EMSI_OFFICIEL.docx"
OUTPUT = ROOT / "rapport" / "RAPPORT_PFE_GROD_FINAL_EMSI.docx"
CAPTURES = ROOT / "rapport" / "captures-finales"
DIAGRAMS = ROOT / "diagrams"

GREEN = "006B3C"
DARK_GREEN = "003C2A"
COPPER = "B75A22"
MID_GRAY = "64706A"
LIGHT_GREEN = "E7F2EC"
LIGHT_COPPER = "F8EAE2"
BLACK = "111111"


def find_paragraph(doc, exact):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == exact:
            return paragraph
    raise ValueError(f"Paragraphe introuvable : {exact}")


def paragraph_text(element):
    return "".join(node.text or "" for node in element.iter(qn("w:t")))


def set_font(run, size=12, bold=None, italic=None, color=BLACK, name="Times New Roman"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(key), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def collect_new(doc, before):
    return [e for e in doc._element.body.iterchildren() if e not in before]


def append_before(doc, anchor, builder):
    before = set(doc._element.body.iterchildren())
    builder()
    for element in collect_new(doc, before):
        anchor._p.addprevious(element)


def add_body_paragraph(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_font(r, bold=True, color=DARK_GREEN)
        r = p.add_run(text[len(bold_prefix):])
        set_font(r)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_heading(doc, text, level=2):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    for run in p.runs:
        set_font(run, size={1: 16, 2: 14, 3: 12}[level], bold=True,
                 color={1: DARK_GREEN, 2: GREEN, 3: COPPER}[level])
    return p


def shade_cell(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=BLACK, size=10):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(doc, caption, headers, rows, widths=None):
    cap = doc.add_paragraph(caption, style="Table Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, color="FFFFFF", size=9.5)
        shade_cell(table.rows[0].cells[i], GREEN)
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            set_cell_text(cells[i], str(value), size=9.5)
            if len(table.rows) % 2 == 1:
                shade_cell(cells[i], "F6F8F7")
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
    doc.add_paragraph()
    return table


def remove_between(doc, start_text, end_text):
    start = find_paragraph(doc, start_text)
    end = find_paragraph(doc, end_text)
    node = start._p.getnext()
    while node is not None and node is not end._p:
        nxt = node.getnext()
        doc._element.body.remove(node)
        node = nxt


def replace_summary_sections(doc):
    remove_between(doc, "Résumé", "Abstract")
    abstract = find_paragraph(doc, "Abstract")

    def french():
        add_body_paragraph(doc,
            "Ce projet de fin d’études porte sur la conception et le développement de G-ROD, une plateforme web B2B destinée à Morocco Copper Foundry. Le besoin initial était double : présenter de manière structurée une gamme de huit produits en cuivre et centraliser le suivi des demandes commerciales et documentaires. La solution réalisée comprend un site public responsive, un catalogue avec recherche, filtres, fiches détaillées et aperçus 3D chargés à la demande, ainsi qu’un espace d’administration consacré aux demandes de devis, aux clients, aux produits, aux ressources techniques et aux notifications.")
        add_body_paragraph(doc,
            "L’architecture s’appuie sur un backend Java 21 avec Spring Boot 3.5.14, une interface React 19.2.6 construite avec Vite 8.2.1 et une base MySQL dans la configuration applicative. La sécurité de l’administration repose sur Spring Security, des jetons JWT, le hachage BCrypt, une limitation locale du débit et un support WebAuthn partiel. Les fichiers téléversés sont contrôlés puis stockés sur disque. La validation réelle a donné 44 tests backend réussis sur 44, un contrôle ESLint réussi, un build Vite réussi et 37 scénarios Playwright réussis sur 38. L’unique échec de recette provient d’images initiales externes devenues injoignables, et non d’un blocage du parcours métier principal.")
        add_body_paragraph(doc,
            "Le résultat est une Release Candidate exploitable pour une démonstration académique et une préparation au déploiement derrière Nginx. L’audit final met toutefois en évidence plusieurs limites : l’absence de chaîne CI/CD, un mode hors ligne PWA incomplet, des modèles 3D génériques qui ne remplacent pas des fichiers CAO, un composant frontend devenu volumineux et des mécanismes de stockage ou de limitation qui restent locaux à une instance. Ces constats servent de base aux perspectives proposées, sans présenter comme livrées des fonctions encore absentes.")
        p = doc.add_paragraph()
        r = p.add_run("Mots-clés : ")
        set_font(r, bold=True, color=DARK_GREEN)
        r = p.add_run("plateforme B2B, cuivre, React, Spring Boot, catalogue industriel.")
        set_font(r)
        doc.add_page_break()

    append_before(doc, abstract, french)

    remove_between(doc, "Abstract", "ملخص")
    arabic = find_paragraph(doc, "ملخص")

    def english():
        add_body_paragraph(doc,
            "This final-year project concerns the design and development of G-ROD, a B2B web platform for Morocco Copper Foundry. The initial need was to present an eight-product copper catalogue in a structured way and to centralize commercial and technical-document requests. The delivered solution combines a responsive public website with search, filters, product detail pages and on-demand 3D previews, and a secured administration area for quotations, customers, products, technical resources and notifications.")
        add_body_paragraph(doc,
            "The implementation uses Java 21 and Spring Boot 3.5.14 for the backend, React 19.2.6 with Vite 8.2.1 for the frontend, and MySQL in the application configuration. Administrative access uses Spring Security, JWT tokens, BCrypt password hashing, local rate limiting and partial WebAuthn support. File uploads are validated and stored on disk. The verified results are 44 successful backend tests out of 44, a successful ESLint check, a successful Vite production build, and 37 successful Playwright scenarios out of 38. The remaining browser-test failure is caused by unreachable external seed-image hosts rather than by the core business workflow.")
        add_body_paragraph(doc,
            "The project therefore provides a demonstrable Release Candidate and the material needed to prepare a reverse-proxy deployment. The final audit also records the current boundaries: no CI/CD workflow, incomplete PWA offline coverage, generic 3D assets rather than CAD models, a large frontend component, and local-only storage and rate-limiting mechanisms. These findings are reported transparently and converted into realistic improvement paths.")
        p = doc.add_paragraph()
        r = p.add_run("Keywords: ")
        set_font(r, bold=True, color=DARK_GREEN)
        r = p.add_run("B2B platform, copper, React, Spring Boot, industrial catalogue.")
        set_font(r)
        doc.add_page_break()

    append_before(doc, arabic, english)

    remove_between(doc, "ملخص", "Table des matières")
    toc = find_paragraph(doc, "Table des matières")

    def arabic_summary():
        texts = [
            "يقدم هذا المشروع منصة G-ROD الموجهة إلى شركة Morocco Copper Foundry. تم إنجازها لتنظيم عرض ثمانية منتجات نحاسية، وتسهيل طلبات عروض الأسعار والوثائق التقنية، ثم جمع متابعتها داخل فضاء إداري واحد. يتضمن الموقع العام كتالوجاً قابلاً للبحث والتصفية، وصفحات مفصلة للمنتجات، ومعاينات ثلاثية الأبعاد عند الطلب، ونماذج لإرسال الطلبات. كما يوفر فضاء الإدارة متابعة الطلبات والعملاء والمنتجات والموارد التقنية والإشعارات.",
            "يعتمد الجزء الخلفي على Java 21 وSpring Boot 3.5.14، بينما تعتمد الواجهة على React 19.2.6 وVite 8.2.1، مع MySQL في إعداد التطبيق. تشمل الحماية Spring Security ورموز JWT وتشفير كلمات المرور بواسطة BCrypt وتحديداً محلياً لمعدل الطلبات، إضافة إلى دعم جزئي لمفاتيح WebAuthn. أسفرت عملية التحقق الفعلية عن نجاح 44 اختباراً خلفياً من أصل 44، ونجاح ESLint وبناء Vite، ونجاح 37 اختبار Playwright من أصل 38. يرتبط الاختبار غير الناجح بصور خارجية لم تعد متاحة، ولا يمنع المسار الوظيفي الأساسي.",
            "أنتج المشروع نسخة مرشحة للإصدار قابلة للعرض والتحضير للنشر خلف Nginx. ويسجل التدقيق النهائي حدوداً واضحة، منها غياب الأتمتة CI/CD، وعدم اكتمال العمل دون اتصال، واعتماد نماذج ثلاثية الأبعاد عامة بدل ملفات هندسية، والحاجة إلى تقسيم الواجهة الكبيرة وتحسين التخزين عند التوسع.",
        ]
        for text in texts:
            p = add_body_paragraph(doc, text)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            ppr = p._p.get_or_add_pPr()
            bidi = OxmlElement("w:bidi")
            bidi.set(qn("w:val"), "1")
            ppr.append(bidi)
            for run in p.runs:
                set_font(run, name="Arial")
        p = doc.add_paragraph("الكلمات المفتاحية: منصة أعمال، النحاس، React، Spring Boot، كتالوج صناعي.")
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in p.runs:
            set_font(run, name="Arial", bold=True, color=DARK_GREEN)
        doc.add_page_break()

    append_before(doc, toc, arabic_summary)


def enrich_introduction(doc):
    anchor = find_paragraph(doc, "1. Contexte général et environnement du projet")

    def builder():
        add_heading(doc, "Problématique", 2)
        add_body_paragraph(doc,
            "Morocco Copper Foundry dispose d’une activité industrielle qui doit être expliquée à des interlocuteurs B2B ayant des attentes précises : nature du produit, usage, pureté, dimensions, norme et disponibilité d’une documentation. Avant la plateforme, ces informations n’étaient pas réunies dans un parcours numérique unique permettant de passer directement de la découverte d’un produit à une demande structurée. Le problème traité n’est donc pas seulement la création d’un site vitrine. Il concerne aussi la qualité des informations reçues par l’équipe commerciale et la capacité à retrouver l’historique d’une société, d’un produit ou d’un document demandé.")
        add_body_paragraph(doc,
            "Une demande imprécise entraîne des échanges supplémentaires avant même qu’une étude commerciale puisse commencer. Inversement, un formulaire trop rigide peut décourager un acheteur qui ne possède pas encore toutes les valeurs techniques. Dans le cadre de ce projet, nous avons cherché un équilibre : guider l’utilisateur avec des champs métier et un indicateur de complétude, tout en conservant des zones libres pour les cas particuliers. L’administration devait ensuite rendre ces informations lisibles, filtrables et prioritaires, y compris lorsque le volume de demandes augmente.")
        add_heading(doc, "Objectifs du projet", 2)
        add_body_paragraph(doc,
            "L’objectif principal consiste à fournir une plateforme B2B cohérente de bout en bout. Pour le visiteur, elle doit présenter l’entreprise, le processus de production, les produits et les ressources disponibles, puis faciliter une prise de contact exploitable. Pour l’administrateur, elle doit centraliser les demandes, proposer une lecture par statut et priorité, retrouver l’historique d’un client et maintenir le contenu utile. Les objectifs techniques associés sont la séparation claire entre frontend et API, la persistance des données, la protection des routes administratives, la validation des fichiers et la préparation d’un déploiement reproductible.")
        add_body_paragraph(doc,
            "Le projet ne cherche pas à remplacer un ERP, à calculer automatiquement un prix industriel ou à produire des plans de fabrication. Il ne contient pas non plus de configurateur CAO. Les aperçus 3D servent à mieux comprendre les familles de produits et restent des représentations génériques. Cette limite de périmètre a été conservée dans le rapport afin que les résultats présentés correspondent exactement à la Release Candidate disponible.")
        add_heading(doc, "Démarche suivie", 2)
        add_body_paragraph(doc,
            "Le travail a été conduit par incréments. Une première étape a porté sur la compréhension du domaine et la définition des flux publics et administratifs. Les pages et les API ont ensuite été développées, puis consolidées par des corrections de navigation, d’affichage responsive, de gestion des images, de présentation des caractéristiques produits et de lisibilité en mode sombre. La dernière phase a regroupé les tests automatisés, la recette des principaux parcours, l’audit de sécurité, la préparation de la configuration de production et la création d’une Release Candidate locale.")
        add_body_paragraph(doc,
            "La validation s’appuie sur des preuves reproductibles : tests Maven sur un profil H2 isolé, analyse ESLint, build de production Vite et scénarios Playwright. Les captures du chapitre 4 ont été produites avec des données fictives dans cet environnement de test. Aucune donnée client réelle n’a été utilisée. Cette précaution répond à la fois à la confidentialité du projet et à la nécessité de documenter des écrans suffisamment remplis pour être compris.")
        add_heading(doc, "Organisation du rapport", 2)
        add_body_paragraph(doc,
            "Le premier chapitre présente l’entreprise, son contexte industriel et le positionnement de G-ROD. Le deuxième formalise le problème, les besoins, le benchmark et la planification. Le troisième justifie la pile technique et décrit l’architecture à l’aide de six diagrammes. Le quatrième relie la réalisation observable aux choix de conception, puis expose les résultats de test, les limites issues de l’audit et les perspectives. Les annexes regroupent enfin les routes, les groupes d’API et les procédures utiles au lancement local.")
        doc.add_page_break()

    append_before(doc, anchor, builder)


def enrich_chapter1(doc):
    anchor = find_paragraph(doc, "Conclusion du chapitre 1")

    def builder():
        add_heading(doc, "1.6 Parties prenantes et usages visés", 2)
        add_body_paragraph(doc,
            "La plateforme met en relation plusieurs usages qui existaient auparavant sous des formes séparées. L’acheteur industriel consulte le catalogue pour vérifier qu’une famille de produit correspond à son besoin. Le bureau d’études s’intéresse davantage aux dimensions, aux normes et aux documents. L’équipe commerciale reçoit la demande et doit identifier rapidement la société, le produit, la quantité et le délai. Enfin, l’administrateur maintient les contenus et suit l’évolution des dossiers. Le modèle retenu ne crée pas de compte client public : le parcours reste volontairement direct et la sécurisation forte est concentrée sur l’espace interne.")
        add_table(doc, "Tableau 1.2 — Parties prenantes et attentes principales",
                  ["Partie prenante", "Attente observée", "Réponse apportée"], [
                      ["Acheteur industriel", "Identifier rapidement un produit et transmettre un besoin", "Catalogue, recherche, fiche produit et devis"],
                      ["Bureau d’études", "Préciser dimensions, norme et plan", "Champs techniques et pièce jointe contrôlée"],
                      ["Équipe commerciale", "Qualifier et prioriser les opportunités", "Pipeline, score, filtres et statuts"],
                      ["Administrateur", "Maintenir produits, ressources et suivi", "Espace protégé et écrans de gestion"],
                  ], [3.5, 5.4, 6.1])
        add_heading(doc, "1.7 Cohérence avec le processus industriel", 2)
        add_body_paragraph(doc,
            "Le schéma de production fourni par l’entreprise donne au site un ancrage concret. Il présente la réception et le tri des matières premières, le contrôle des pesées, le stockage, la transformation, le laboratoire, la conformité et l’expédition. La plateforme ne pilote aucun de ces équipements. Elle reprend seulement cette lecture pour expliquer, à un visiteur non spécialiste, comment une matière est contrôlée puis transformée avant d’être proposée sous différentes formes. Cette distinction évite de confondre information institutionnelle et système de supervision industrielle.")
        add_body_paragraph(doc,
            "Les produits du catalogue correspondent aux familles manipulées dans les échanges B2B : rod, anodes, bus bars, flat bars, tubes, sheets, wire et pièces personnalisées. Les informations affichées ont été structurées autour de la pureté, des dimensions et des normes. Lorsque la valeur exacte dépend du cahier des charges, le texte le signale au lieu de présenter une mesure arbitraire. Ce choix a notamment permis de supprimer les zones vides observées dans les premières versions tout en restant prudent sur les données techniques non validées par l’entreprise.")
        add_heading(doc, "1.8 Valeur apportée par la centralisation", 2)
        add_body_paragraph(doc,
            "Le bénéfice le plus visible est la continuité du parcours. Une personne peut partir d’une recherche, ouvrir une fiche, visualiser une géométrie, consulter une ressource ou envoyer une demande préremplie avec le produit choisi. Côté administration, les mêmes informations alimentent les compteurs, la liste, le pipeline et l’historique client. Il n’est donc pas nécessaire de ressaisir manuellement le nom du produit pour rapprocher les écrans. Cette continuité provient de l’usage d’une API commune et de références métier partagées entre les pages.")
        add_body_paragraph(doc,
            "La centralisation ne signifie pas que tous les processus de l’entreprise sont automatisés. Le chiffrage reste une décision humaine et les documents restreints sont transmis selon le traitement retenu par l’équipe. La plateforme aide à préparer et suivre le dossier ; elle ne remplace pas la validation technique, commerciale ou qualité. Cette frontière est importante pour apprécier correctement le résultat du projet.")

    append_before(doc, anchor, builder)


def renumber_chapter2(doc):
    replacements = {
        "2.3 Acteurs du système": "2.4 Acteurs du système",
        "2.4 Besoins fonctionnels": "2.5 Besoins fonctionnels",
        "2.4.1 Site public": "2.5.1 Site public",
        "2.4.2 Espace administrateur": "2.5.2 Espace administrateur",
        "2.5 Exigences non fonctionnelles": "2.6 Exigences non fonctionnelles",
        "2.6 Périmètre du projet": "2.7 Périmètre du projet",
        "2.7 Méthodologie de travail": "2.8 Méthodologie de travail",
        "2.8 Planification": "2.9 Planification",
        "2.9 Gestion des risques et confidentialité": "2.10 Gestion des risques et confidentialité",
        "Tableau 2.2 — Besoins fonctionnels du site public": "Tableau 2.3 — Besoins fonctionnels du site public",
        "Tableau 2.3 — Besoins fonctionnels de l'administration": "Tableau 2.4 — Besoins fonctionnels de l'administration",
        "Tableau 2.4 — Exigences non fonctionnelles": "Tableau 2.5 — Exigences non fonctionnelles",
        "Tableau 2.5 — Planification par phases": "Tableau 2.6 — Planification par phases",
        "Tableau 2.6 — Principaux risques du projet": "Tableau 2.7 — Principaux risques du projet",
    }
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text in replacements:
            paragraph.text = replacements[text]
            if text.startswith("Tableau"):
                paragraph.style = "Table Caption"
            elif re.match(r"^2\.\d+\.\d+", replacements[text]):
                paragraph.style = "Heading 3"
            else:
                paragraph.style = "Heading 2"


def add_benchmark(doc):
    renumber_chapter2(doc)
    anchor = find_paragraph(doc, "2.4 Acteurs du système")

    def builder():
        add_heading(doc, "2.3 Benchmark de plateformes industrielles B2B", 2)
        add_heading(doc, "2.3.1 Objectif et méthode", 3)
        add_body_paragraph(doc,
            "Le benchmark a été réalisé sur quatre acteurs industriels disposant de contenus publics officiels : Wieland, KME, Aurubis et Nexans. L’objectif n’était pas de reproduire leurs portails ni de comparer des fonctions internes non accessibles. L’analyse s’est limitée à ce qu’un acheteur peut réellement observer : organisation du catalogue, accès à la documentation, niveau de détail des fiches, moteurs de recherche et mécanismes de prise de contact. Les constats sont donc utilisés comme références d’ergonomie et non comme preuve sur leurs systèmes d’information internes.")
        add_heading(doc, "2.3.2 Wieland", 3)
        add_body_paragraph(doc,
            "Wieland organise son offre autour des produits et de leurs domaines d’application. La navigation donne une place importante aux formes de cuivre, aux alliages et aux usages, ce qui aide un visiteur à raisonner soit par matière, soit par besoin [9]. Pour G-ROD, cette observation a confirmé l’intérêt de conserver à la fois un catalogue explicite et un assistant de recherche par application. La plateforme réalisée reste plus limitée : elle couvre huit familles et ne propose pas de sélection métallurgique détaillée. Cette différence est cohérente avec son périmètre de projet de fin d’études.")
        add_heading(doc, "2.3.3 KME", 3)
        add_body_paragraph(doc,
            "Le site KME met en avant un centre de téléchargement qui regroupe brochures, fiches et documents associés aux activités industrielles [10]. Cette organisation montre qu’un catalogue seul ne suffit pas dans un contexte B2B : les acheteurs recherchent aussi des éléments de conformité et de préparation technique. G-ROD reprend cette idée sous une forme adaptée, avec une page publique de ressources, des filtres par type et un formulaire distinct lorsqu’un document doit être demandé. L’administration peut ajouter une ressource, la lier à un produit et choisir si son téléchargement est public ou restreint.")
        add_heading(doc, "2.3.4 Aurubis", 3)
        add_body_paragraph(doc,
            "Aurubis présente plusieurs produits en cuivre et met également en avant des services destinés à ses partenaires [11]. Cette combinaison entre information industrielle et relation commerciale a été retenue comme un repère pour la page d’accueil de G-ROD. Le site ne se limite pas à afficher des images : il propose des points d’entrée directs vers le catalogue, les documents et le devis. En revanche, aucun portail partenaire avancé n’a été développé. Les visiteurs publics n’ont pas de compte et le suivi détaillé reste réservé à l’administrateur.")
        add_heading(doc, "2.3.5 Nexans", 3)
        add_body_paragraph(doc,
            "Le catalogue Nexans illustre une recherche structurée dans une offre technique étendue, avec des catégories et des informations destinées à guider la sélection [12]. Dans G-ROD, le volume est nettement plus faible, mais la même logique de réduction de l’effort de recherche apparaît à travers le champ texte, le filtre par application, le tri, le choix entre cartes et liste et l’affichage des produits actifs. L’interface recommande également un produit à partir des critères disponibles. Il s’agit d’une aide déterministe fondée sur les données du catalogue, et non d’un moteur d’intelligence artificielle entraîné.")
        add_table(doc, "Tableau 2.2 — Synthèse du benchmark industriel B2B",
                  ["Solution", "Élément observé", "Apport retenu pour G-ROD", "Limite de comparaison"], [
                      ["Wieland", "Produits et applications", "Double entrée catalogue/application", "Gamme et alliages beaucoup plus étendus"],
                      ["KME", "Centre de téléchargement", "Ressources filtrables et liées aux produits", "Documents internes non observables"],
                      ["Aurubis", "Produits et relation partenaires", "Passerelles vers demandes commerciales", "Pas de portail partenaire livré"],
                      ["Nexans", "Catalogue technique structuré", "Recherche, filtres, tri et fiches", "Secteur et profondeur de catalogue différents"],
                  ], [2.4, 3.7, 5.0, 3.9])
        add_heading(doc, "2.3.6 Enseignements pour le cadrage", 3)
        add_body_paragraph(doc,
            "Trois enseignements ont guidé le cadrage. Premièrement, la documentation doit rester accessible depuis le produit et depuis un espace dédié. Deuxièmement, les filtres n’ont de valeur que si les données du catalogue sont réellement renseignées. Troisièmement, un appel à l’action commercial doit rester visible sans interrompre la consultation. Ces principes se retrouvent dans les parcours présentés au chapitre 4. Le benchmark n’a pas conduit à ajouter des fonctions étrangères au besoin, telles qu’un paiement en ligne, un calcul tarifaire automatique ou un extranet client complet.")

    append_before(doc, anchor, builder)


def add_technical_comparisons(doc):
    anchor = find_paragraph(doc, "3.3 Architecture logique")

    def builder():
        add_heading(doc, "3.2.1 Comparaison des solutions frontend", 3)
        add_body_paragraph(doc,
            "Le frontend devait gérer des pages publiques, des formulaires riches, une administration interactive et un visualiseur 3D. React a été retenu parce que le projet pouvait représenter ces fonctions par des composants partageant le même état et les mêmes services d’accès à l’API [2]. Angular aurait fourni un cadre plus intégré, mais avec une structure et un apprentissage plus lourds pour le périmètre. Vue constituait aussi une option valable ; le choix final a surtout été guidé par la maîtrise de React et son intégration directe avec React Router, Three.js et l’outillage Vite.")
        add_table(doc, "Tableau 3.2 — Comparaison des solutions frontend",
                  ["Critère", "React", "Angular", "Vue"], [
                      ["Organisation", "Bibliothèque composable", "Framework complet", "Framework progressif"],
                      ["Prise en main dans le projet", "Maîtrisée", "Plus lourde", "Accessible"],
                      ["Écosystème 3D et tests", "Intégration directe", "Possible", "Possible"],
                      ["Décision", "Retenu", "Non retenu", "Non retenu"],
                  ], [3.7, 4.0, 4.0, 4.0])
        add_heading(doc, "3.2.2 Comparaison des solutions backend", 3)
        add_body_paragraph(doc,
            "Le backend porte les règles de validation, la persistance, les fichiers, les notifications et l’authentification. Spring Boot fournit un cadre cohérent pour les contrôleurs REST, Spring Data JPA, la validation et Spring Security [1]. Express aurait permis une API légère en JavaScript, tandis que Django propose un environnement Python très productif. Spring Boot a toutefois été retenu pour son typage, l’intégration des composants de sécurité et la facilité de structurer des tests d’intégration avec JUnit, MockMvc et H2.")
        add_table(doc, "Tableau 3.3 — Comparaison des solutions backend",
                  ["Critère", "Spring Boot", "Express", "Django"], [
                      ["Typage et structure", "Java typé, conventions fortes", "Souple, structure à définir", "Python, cadre intégré"],
                      ["Sécurité", "Spring Security", "Bibliothèques à assembler", "Mécanismes intégrés"],
                      ["Persistance", "Spring Data JPA", "ORM au choix", "Django ORM"],
                      ["Décision", "Retenu", "Non retenu", "Non retenu"],
                  ], [3.7, 4.0, 4.0, 4.0])
        add_heading(doc, "3.2.3 Comparaison des bases de données", 3)
        add_body_paragraph(doc,
            "Les données principales sont relationnelles : utilisateurs, clients, demandes, produits, documents et lectures de notifications. MySQL a été choisi pour la configuration de l’application et son usage courant dans un environnement XAMPP [3]. PostgreSQL aurait également convenu et offre des fonctions avancées, mais n’apportait pas d’avantage décisif pour le modèle actuel. MongoDB aurait simplifié certains objets variables, au prix d’une cohérence relationnelle moins naturelle pour le suivi commercial. H2 n’est pas la base de production : il sert uniquement à isoler et accélérer les tests automatisés.")
        add_table(doc, "Tableau 3.4 — Comparaison des bases de données",
                  ["Critère", "MySQL", "PostgreSQL", "MongoDB"], [
                      ["Modèle principal", "Relationnel", "Relationnel avancé", "Document"],
                      ["Adéquation aux entités", "Bonne", "Bonne", "Moins directe"],
                      ["Environnement du projet", "Disponible avec XAMPP", "Installation distincte", "Installation distincte"],
                      ["Décision", "Retenu", "Non retenu", "Non retenu"],
                  ], [3.7, 4.0, 4.0, 4.0])
        add_body_paragraph(doc,
            "Vite complète ce choix en assurant le serveur de développement et le build du frontend [4]. Three.js est chargé dans un morceau séparé pour les scènes 3D [5], et Playwright automatise les parcours navigateur [6]. Les recommandations OWASP ont servi de repère pour les en-têtes, la validation et l’authentification [7], tandis que l’API WebAuthn est utilisée pour le mécanisme optionnel de Passkey [8]. Ces références décrivent les technologies ; la validation de leur usage concret dans G-ROD provient du code et des tests du dépôt [14].")

    append_before(doc, anchor, builder)

    for paragraph in doc.paragraphs:
        replacements = {
            "Tableau 3.2 — Principaux cas d'utilisation": "Tableau 3.5 — Principaux cas d'utilisation",
            "Tableau 3.3 — Entités métier": "Tableau 3.6 — Entités métier",
        }
        if paragraph.text.strip() in replacements:
            paragraph.text = replacements[paragraph.text.strip()]
            paragraph.style = "Table Caption"


def add_design_analysis(doc):
    anchor = find_paragraph(doc, "Conclusion du chapitre 3")

    def builder():
        add_heading(doc, "3.9 Cohérence des diagrammes avec l’implémentation", 2)
        add_body_paragraph(doc,
            "Les six diagrammes de ce chapitre ont été reconstruits à partir des routes, contrôleurs, services et entités présents dans le dépôt. L’architecture globale distingue le navigateur, le frontend React, l’API Spring Boot, MySQL et le stockage sur disque. Cette représentation ne fait apparaître ni microservice ni service cloud, car la Release Candidate est une application déployable comme un ensemble frontend/backend. Nginx est présenté comme cible de reverse proxy dans les fichiers de préparation, et non comme un déploiement déjà réalisé.")
        add_body_paragraph(doc,
            "Le diagramme de cas d’utilisation rassemble les fonctions publiques et administratives. Il montre que le visiteur peut consulter, rechercher, visualiser et envoyer des demandes, alors que l’administrateur doit s’authentifier avant d’accéder aux opérations de suivi et de maintenance. Le formulaire de devis n’exige pas la création préalable d’un compte client. À la réception, le backend crée ou rapproche le client à partir des informations disponibles puis enregistre la demande. Cette organisation correspond à l’objectif de réduire les obstacles au premier contact.")
        add_body_paragraph(doc,
            "Les trois diagrammes de séquence détaillent les points sensibles. Pour un devis, la validation du formulaire et du fichier précède l’écriture métier. Lorsqu’un fichier temporaire doit devenir définitif, sa promotion sur disque intervient au début du traitement transactionnel, avant les écritures en base. Cette séquence explique une limite relevée par l’audit : un échec de base après la promotion pourrait laisser un fichier orphelin. Pour l’authentification, la vérification des identifiants produit un JWT utilisé ensuite dans l’en-tête Authorization. Pour le traitement administratif, le changement de statut déclenche la persistance et la création d’une notification consultable dans le centre dédié.")
        add_body_paragraph(doc,
            "Le modèle JPA contient dix classes annotées comme entités. Seules trois associations sont matérialisées par des relations JPA : DemandeDevis vers Client, puis NotificationLecture vers Notification et Utilisateur. Les autres références métier sont stockées sous forme de valeurs scalaires. Le diagramme le montre volontairement, au lieu de supposer des associations qui n’existent pas. Ce choix facilite la lecture de l’état actuel et permet d’identifier les normalisations possibles comme une perspective, sans les confondre avec le travail livré.")
        add_heading(doc, "3.10 Choix de sécurité et compromis", 2)
        add_body_paragraph(doc,
            "La protection de l’administration combine plusieurs niveaux. Spring Security filtre les routes, le mot de passe est haché par BCrypt et le JWT transporte l’identité pour les appels suivants. Des en-têtes de sécurité et des règles CORS configurables complètent ce mécanisme. Une limitation de débit réduit les tentatives répétées sur certains points d’entrée. Son stockage est cependant local à la mémoire de l’instance : il devra être externalisé si plusieurs serveurs partagent le trafic.")
        add_body_paragraph(doc,
            "Le support Passkey repose sur WebAuthn, mais l’audit ne le classe pas comme totalement abouti. Les challenges sont conservés en mémoire, le compteur de signature n’est pas mis à jour et certains indicateurs utilisateur ne sont pas explicitement vérifiés. Dans le rapport, cette fonction est donc décrite comme un support partiel et optionnel. Le mode de connexion JWT reste le mécanisme principal validé par les scénarios de recette.")

    append_before(doc, anchor, builder)


def fix_known_inaccuracies(doc):
    replacements = {
        "React Router 7.15.1": "React Router 7.18.2",
        "Vite 8.0.12": "Vite 8.2.1",
        "38/38": "37/38",
        "v1.0.0-rc.2": "v1.0.0-rc.1",
        "/ressources": "/resources",
    }
    for paragraph in doc.paragraphs:
        for old, new in replacements.items():
            if old in paragraph.text:
                for run in paragraph.runs:
                    if old in run.text:
                        run.text = run.text.replace(old, new)
                if old in paragraph.text:
                    paragraph.text = paragraph.text.replace(old, new)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for old, new in replacements.items():
                    if old in cell.text:
                        cell.text = cell.text.replace(old, new)


def replace_entity_table(doc):
    target = None
    for table in doc.tables:
        if table.rows and table.rows[0].cells[0].text.strip() == "Entité":
            target = table
            break
    if target is None:
        return
    while len(target.rows) > 1:
        target._tbl.remove(target.rows[-1]._tr)
    rows = [
        ("Utilisateur", "Compte administrateur et informations de sécurité"),
        ("WebAuthnCredential", "Identifiant et clé publique d’une Passkey"),
        ("Produit", "Catalogue, caractéristiques, image et état actif"),
        ("Client", "Société, contact et informations de suivi"),
        ("DemandeDevis", "Besoin commercial lié par ManyToOne à Client"),
        ("DocumentRequest", "Demande publique de document technique"),
        ("TechnicalResource", "Ressource téléchargeable, visibilité et produit concerné"),
        ("Notification", "Événement généré pour l’administration"),
        ("NotificationLecture", "Lecture liée à Notification et Utilisateur"),
        ("FileAudit", "Trace de contrôle d’un fichier stocké"),
    ]
    for name, role in rows:
        cells = target.add_row().cells
        set_cell_text(cells[0], name, bold=True, color=DARK_GREEN)
        set_cell_text(cells[1], role)


def enrich_tests_and_limits(doc):
    anchor = find_paragraph(doc, "4.7 Préparation du déploiement")

    def tests():
        add_heading(doc, "4.6.1 Stratégie de validation", 3)
        add_body_paragraph(doc,
            "La validation a été organisée en couches afin de distinguer les erreurs de logique backend, les défauts statiques du frontend, les problèmes de compilation et les ruptures de parcours dans le navigateur. Le backend est lancé avec un profil de test H2, ce qui évite de modifier une base MySQL de travail. Les scénarios Playwright utilisent également des comptes et des données fictives. Cette séparation rend les essais reproductibles et protège les informations métier.")
        add_heading(doc, "4.6.2 Tests backend : 44/44", 3)
        add_body_paragraph(doc,
            "La commande Maven de test termine avec 44 tests réussis sur 44, sans échec ni erreur. Les contrôles couvrent notamment l’authentification, les demandes, les validations, les fichiers et plusieurs règles de sécurité. Ce résultat montre que les cas automatisés existants sont stables dans l’environnement H2. Il ne prouve pas à lui seul le comportement sur une base MySQL de production ni la résistance sous forte charge ; ces deux aspects restent explicitement hors de la recette effectuée.")
        add_heading(doc, "4.6.3 Qualité statique et build frontend", 3)
        add_body_paragraph(doc,
            "ESLint s’exécute avec succès sur le code frontend. Le build Vite de production aboutit également et génère les ressources distribuables. La sortie met toutefois en évidence un morceau JavaScript lié à la 3D d’environ 614,5 kB, supérieur au seuil d’avertissement de 500 kB. Le chargement différé du visualiseur limite déjà son impact sur le premier affichage, mais un découpage complémentaire et une vérification de la compression restent conseillés avant une diffusion à grande échelle.")
        add_heading(doc, "4.6.4 Recette Playwright : 37/38", 3)
        add_body_paragraph(doc,
            "La campagne Playwright comporte 38 scénarios, dont 37 réussissent. Les cinq parcours identifiés comme critiques passent, ainsi que 32 autres scénarios. Ils couvrent notamment la navigation publique, le catalogue, les demandes, la connexion administrateur et les écrans de suivi. L’unique échec est causé par huit URL d’images initiales pointant vers des domaines externes qui ne sont plus résolus ; le navigateur remonte net::ERR_NAME_NOT_RESOLVED. Les images locales ajoutées par l’administration fonctionnent et le défaut ne bloque pas les opérations métier principales.")
        add_heading(doc, "4.6.5 Lecture critique des résultats", 3)
        add_body_paragraph(doc,
            "Le bilan n’est donc pas présenté comme un résultat parfait. La Release Candidate satisfait les contrôles principaux, mais le jeu de données initial doit cesser de dépendre d’hôtes externes pour obtenir une recette complètement verte. La solution recommandée est de remplacer ces URL par des ressources locales versionnées ou gérées par le stockage applicatif. Cette correction concerne les données de démonstration ; elle ne nécessite pas de modifier le fonctionnement du catalogue.")
        add_table(doc, "Tableau 4.3 — Résultats détaillés de la validation finale",
                  ["Contrôle", "Résultat réel", "Interprétation"], [
                      ["Backend Maven", "44/44 réussis", "Aucun échec dans le profil H2"],
                      ["ESLint", "Réussi", "Aucune erreur bloquante signalée"],
                      ["Build Vite", "Réussi", "Artefacts produits ; avertissement chunk 3D"],
                      ["Playwright", "37/38 réussis", "Échec lié aux images seed externes"],
                      ["Parcours critiques", "5/5 réussis", "Flux principaux utilisables"],
                  ], [3.5, 3.8, 7.7])

    append_before(doc, anchor, tests)

    anchor2 = find_paragraph(doc, "Conclusion du chapitre 4")

    def limits():
        add_heading(doc, "4.8.1 Résultats obtenus", 3)
        add_body_paragraph(doc,
            "La version finale réunit le site public et l’administration autour d’une même API. Les huit produits sont consultables avec leurs caractéristiques, leur image et leur aperçu 3D lorsqu’il est disponible. Les demandes de devis et de documents sont persistées, recherchées, filtrées et affichées dans des tableaux paginés. Le pipeline limite le nombre de cartes prioritaires visibles dans chaque colonne et propose l’accès aux demandes restantes, ce qui évite l’allongement illimité de l’écran lorsque le volume augmente. Les ressources techniques, les notifications et l’historique client complètent ce suivi.")
        add_body_paragraph(doc,
            "Les corrections d’interface réalisées pendant la consolidation concernent notamment la mise en page du mode liste, l’alignement des filtres, la lisibilité du mode sombre, les colonnes d’actions, les formulaires de ressources et l’ajout d’images produits. Les captures de ce chapitre montrent le résultat obtenu dans un environnement de recette rempli avec des données fictives. Elles ne constituent pas une preuve de déploiement public.")
        add_heading(doc, "4.8.2 Limites confirmées par l’audit", 3)
        add_body_paragraph(doc,
            "Le frontend concentre encore une part importante de la logique dans App.jsx, qui dépasse six mille lignes, tandis que la feuille App.css dépasse quatorze mille lignes. Cette taille complique la maintenance et les tests ciblés. Une décomposition par domaine — catalogue, devis, administration, notifications et compte — réduirait le couplage. Il s’agit d’une amélioration de structure ; aucune réécriture de ce type n’a été réalisée uniquement pour le rapport.")
        add_body_paragraph(doc,
            "La PWA est installable grâce au manifeste et au service worker, mais son mode hors ligne est incomplet. La route racine n’est pas préchargée et le modèle GLB des flat bars n’est pas couvert par la liste de précache. La plateforme ne doit donc pas être présentée comme entièrement utilisable sans réseau. De la même manière, les huit modèles 3D sont des géométries génériques adaptées à la visualisation. Ils ne remplacent pas des plans CAO certifiés et l’administration ne prend pas encore en charge l’import de fichiers GLB ou GLTF.")
        add_body_paragraph(doc,
            "Le stockage des téléversements s’effectue sur le disque local et la limitation de débit reste en mémoire. Ces choix conviennent à une instance de démonstration ou à un premier déploiement maîtrisé, mais ils ne sont pas partagés entre plusieurs serveurs. La promotion d’un fichier temporaire intervient avant la fin des écritures transactionnelles en base ; une erreur ultérieure peut donc laisser un fichier final orphelin. Le service d’audit sait repérer des incohérences, mais ne réalise pas automatiquement leur nettoyage.")
        add_body_paragraph(doc,
            "L’envoi de courriels, les notifications SMS et l’assistant conversationnel dépendent d’une configuration externe. Ils peuvent être désactivés et ne doivent pas être considérés comme garantis sans clés ni fournisseurs. Le support WebAuthn est partiel : les challenges sont conservés en mémoire, le compteur de signature n’est pas actualisé et certaines vérifications pourraient être renforcées. Enfin, aucune chaîne CI/CD n’est présente dans le dépôt, aucun déploiement public n’a été vérifié et aucune restauration complète de sauvegarde MySQL n’a fait partie de l’audit.")
        add_heading(doc, "4.8.3 Perspectives réalistes", 3)
        add_body_paragraph(doc,
            "La première perspective consiste à stabiliser totalement la recette en remplaçant les images initiales externes par des fichiers locaux contrôlés. La deuxième est de modulariser le frontend sans modifier les parcours, puis d’ajouter des tests ciblés sur les composants extraits. La troisième concerne l’exploitation : intégrer une chaîne CI/CD qui exécute Maven, ESLint, le build et Playwright, puis produit des artefacts identifiables avant tout déploiement.")
        add_body_paragraph(doc,
            "Pour une montée en charge horizontale, les fichiers devraient être placés dans un stockage partagé ou objet et la limitation de débit dans un service commun tel que Redis. Le nettoyage compensatoire des fichiers doit également être renforcé lorsque la transaction métier échoue. Le mode PWA pourrait précacher les ressources indispensables et prévoir une page hors ligne dédiée. Côté sécurité, la finalisation WebAuthn nécessiterait un stockage distribué des challenges, la mise à jour du signCount et une validation plus stricte des indicateurs d’authentification.")
        add_body_paragraph(doc,
            "À plus long terme, des modèles 3D issus des données techniques validées pourraient remplacer les formes génériques, avec un véritable processus d’import et de contrôle. Une intégration ERP ou Odoo pourrait éviter certaines doubles saisies, mais elle suppose une étude fonctionnelle et des API d’entreprise qui ne faisaient pas partie du projet. Ces pistes sont présentées comme des prolongements et non comme des fonctionnalités déjà disponibles.")

    append_before(doc, anchor2, limits)


def fit_width(image_path, max_width=14.8, max_height=17.0):
    with Image.open(image_path) as img:
        ratio = img.width / img.height
    return min(max_width, max_height * ratio)


def clear_cell(cell):
    tc = cell._tc
    for child in list(tc):
        if child.tag != qn("w:tcPr"):
            tc.remove(child)


def put_images(cell, image_paths):
    clear_cell(cell)
    count = len(image_paths)
    max_height = 8.0 if count > 1 else 17.0
    for image_path in image_paths:
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        width = fit_width(image_path, 14.8, max_height)
        p.add_run().add_picture(str(image_path), width=Cm(width))
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def put_test_summary(cell):
    clear_cell(cell)
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("VALIDATION AUTOMATISÉE DE LA RELEASE CANDIDATE")
    set_font(r, size=13, bold=True, color=DARK_GREEN)
    nested = cell.add_table(rows=1, cols=3)
    nested.style = "Table Grid"
    for i, value in enumerate(("Contrôle", "Résultat", "État")):
        set_cell_text(nested.rows[0].cells[i], value, bold=True, color="FFFFFF")
        shade_cell(nested.rows[0].cells[i], GREEN)
    for control, result, state in [
        ("Backend", "44/44", "RÉUSSI"),
        ("ESLint", "0 erreur bloquante", "RÉUSSI"),
        ("Build Vite", "Artefacts générés", "RÉUSSI"),
        ("Playwright", "37/38", "1 ÉCART DOCUMENTÉ"),
        ("Parcours critiques", "5/5", "RÉUSSI"),
    ]:
        cells = nested.add_row().cells
        for i, value in enumerate((control, result, state)):
            set_cell_text(cells[i], value, bold=i == 2, color=GREEN if state == "RÉUSSI" else COPPER)
    p = cell.add_paragraph("Écart Playwright : huit images de données initiales utilisent des hôtes externes non résolus.")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        set_font(run, size=9, italic=True, color=MID_GRAY)


def next_caption(table):
    node = table._tbl.getnext()
    while node is not None:
        if node.tag == qn("w:p"):
            text = paragraph_text(node).strip()
            if text:
                return Paragraph(node, table._parent)
        node = node.getnext()
    return None


def previous_text_paragraph(table):
    node = table._tbl.getprevious()
    while node is not None:
        if node.tag == qn("w:p") and paragraph_text(node).strip():
            return Paragraph(node, table._parent)
        node = node.getprevious()
    return None


def replace_placeholders(doc):
    diagram_map = {
        "Figure 2.1": ([DIAGRAMS / "07_gantt_mars_aout_2026.png"], "La figure 2.1 synthétise la planification de mars à août 2026. Les phases antérieures au premier historique Git vérifiable sont signalées comme périodes à confirmer ; les travaux de consolidation et la Release Candidate locale sont, eux, appuyés par les traces du dépôt."),
        "Figure 3.1": ([DIAGRAMS / "01_architecture_globale.png"], "La figure 3.1 présente l’architecture réellement préparée : navigateur, frontend React servi derrière Nginx, API Spring Boot, base MySQL et stockage documentaire sur disque. Elle n’introduit aucun microservice absent du dépôt."),
        "Figure 3.2": ([DIAGRAMS / "02_cas_utilisation_global.png"], "La figure 3.2 regroupe les cas d’utilisation publics et administratifs. Elle met en évidence l’absence de compte client public et la protection des fonctions de gestion par authentification."),
        "Figure 3.3": ([DIAGRAMS / "03_sequence_creation_devis.png"], "La figure 3.3 suit la création d’un devis depuis React jusqu’aux contrôles, au stockage éventuel du fichier, au client et à la demande persistée."),
        "Figure 3.4": ([DIAGRAMS / "04_sequence_authentification_jwt.png"], "La figure 3.4 détaille la connexion principale par mot de passe et JWT. WebAuthn reste un mécanisme optionnel séparé dont les limites sont exposées dans l’audit."),
        "Figure 3.5": ([DIAGRAMS / "05_sequence_traitement_demande.png"], "La figure 3.5 montre la consultation puis la mise à jour d’une demande dans l’administration, avec persistance du statut et génération de notification."),
        "Figure 3.6": ([DIAGRAMS / "06_modele_metier_jpa.png"], "La figure 3.6 représente les dix entités détectées dans le code et uniquement les associations JPA explicites, afin de ne pas inventer de relations."),
        "Figure 4.3": ([CAPTURES / "fig-4-01-accueil-public.png"], "La figure 4.3 montre la page d’accueil publique, ses accès directs au catalogue et au devis, ainsi que l’identité industrielle retenue."),
        "Figure 4.4": ([CAPTURES / "fig-4-02-catalogue-recherche-cartes.png", CAPTURES / "fig-4-03-catalogue-mode-liste.png"], "La figure 4.4 rapproche la recherche du catalogue et son affichage en liste. Les mêmes huit produits restent disponibles dans les deux modes."),
        "Figure 4.5": ([CAPTURES / "fig-4-04-fiche-produit.png", CAPTURES / "fig-4-05-apercu-3d.png"], "La figure 4.5 associe la fiche Copper Rod à son aperçu 3D chargé à la demande. Les caractéristiques affichées proviennent des données produit."),
        "Figure 4.6": ([CAPTURES / "fig-4-06-formulaire-devis.png"], "La figure 4.6 présente le configurateur de devis, l’indicateur de complétude et les informations de contact fictives utilisées pour la recette."),
        "Figure 4.7": ([CAPTURES / "fig-4-07-ressources-publiques.png"], "La figure 4.7 montre l’espace public de ressources avec recherche, filtre par type et distinction entre disponibilité directe et document sur demande."),
        "Figure 4.8": ([CAPTURES / "fig-4-09-catalogue-mobile.png"], "La figure 4.8 vérifie l’adaptation du catalogue à une largeur mobile. La navigation et les filtres passent sur plusieurs lignes sans masquer le contenu."),
        "Figure 4.9": ([CAPTURES / "fig-4-11-dashboard-admin.png"], "La figure 4.9 présente le tableau de bord alimenté par les données de recette : compteurs, demandes récentes, état du catalogue et actions rapides."),
        "Figure 4.10": ([CAPTURES / "fig-4-12-pipeline-commercial.png"], "La figure 4.10 montre la répartition des demandes par statut. Chaque colonne affiche les cartes les plus prioritaires et garde un accès au reste de la liste."),
        "Figure 4.11": ([CAPTURES / "fig-4-13-demandes-admin.png"], "La figure 4.11 illustre la table détaillée des demandes, avec score, statut et actions. Les sociétés et adresses visibles sont fictives."),
        "Figure 4.12": ([CAPTURES / "fig-4-14-clients-historique.png"], "La figure 4.12 présente les clients consolidés à partir des demandes et l’accès à leur historique."),
        "Figure 4.13": ([CAPTURES / "fig-4-15-produits-admin.png"], "La figure 4.13 montre les huit produits, leurs caractéristiques, leur état et la disponibilité d’un modèle 3D générique."),
        "Figure 4.14": ([CAPTURES / "fig-4-14b-demandes-documents.png"], "La figure 4.14 expose les demandes de documents techniques avec recherche, filtres, statuts et accès au détail."),
        "Figure 4.15": ([CAPTURES / "fig-4-16-ressources-admin.png"], "La figure 4.15 montre les ressources ajoutées, leur type, le produit concerné, leur visibilité et les actions disponibles."),
        "Figure 4.16": ([CAPTURES / "fig-4-17-notifications-admin.png"], "La figure 4.16 présente les notifications persistantes créées par les événements de devis et de documents."),
        "Figure 4.17": ([CAPTURES / "fig-4-18-compte-securite.png"], "La figure 4.17 rassemble le changement de mot de passe et l’activation optionnelle d’une Passkey dans la page de compte."),
        "Figure 4.18": ([CAPTURES / "fig-4-18b-navigation-admin-mobile.png"], "La figure 4.18 vérifie le comportement du tiroir de navigation administrateur sur un écran mobile."),
        "Figure 4.19": ([CAPTURES / "fig-4-10-connexion-admin.png"], "La figure 4.19 montre l’écran de connexion administrateur. Le mot de passe est masqué et aucune valeur secrète n’est révélée dans le rapport."),
        "Figure 4.20": (None, "La figure 4.20 synthétise les résultats effectivement obtenus lors de la validation de la Release Candidate."),
    }
    caption_replacements = {
        "Figure 3.2": "Figure 3.2 — Diagramme global des cas d’utilisation",
        "Figure 3.3": "Figure 3.3 — Diagramme de séquence d’une demande de devis",
        "Figure 3.4": "Figure 3.4 — Diagramme de séquence de l’authentification JWT",
        "Figure 3.5": "Figure 3.5 — Diagramme de séquence du traitement administratif",
        "Figure 3.6": "Figure 3.6 — Modèle métier JPA constaté dans le code",
        "Figure 4.4": "Figure 4.4 — Catalogue : recherche et mode liste",
        "Figure 4.5": "Figure 4.5 — Fiche produit et aperçu 3D",
        "Figure 4.7": "Figure 4.7 — Ressources techniques publiques",
        "Figure 4.8": "Figure 4.8 — Catalogue en affichage mobile",
        "Figure 4.14": "Figure 4.14 — Gestion des demandes de documents",
        "Figure 4.18": "Figure 4.18 — Navigation Admin sur écran mobile",
        "Figure 4.19": "Figure 4.19 — Connexion sécurisée à l’administration",
        "Figure 4.20": "Figure 4.20 — Synthèse des tests automatisés",
    }
    for table in doc.tables:
        if not any("ZONE RÉSERVÉE" in cell.text for row in table.rows for cell in row.cells):
            continue
        caption = next_caption(table)
        if caption is None:
            continue
        key = next((prefix for prefix in diagram_map if caption.text.strip().startswith(prefix)), None)
        if key is None:
            continue
        paths, analysis = diagram_map[key]
        if paths:
            for path in paths:
                if not path.exists():
                    raise FileNotFoundError(path)
            put_images(table.cell(0, 0), paths)
        else:
            put_test_summary(table.cell(0, 0))
        row = table.rows[0]
        row.height_rule = WD_ROW_HEIGHT_RULE.AUTO
        for tr_height in list(row._tr.xpath("./w:trPr/w:trHeight")):
            tr_height.getparent().remove(tr_height)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        previous = previous_text_paragraph(table)
        if previous is not None:
            previous.text = analysis
            previous.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for run in previous.runs:
                set_font(run)
        if key in caption_replacements:
            caption.text = caption_replacements[key]
        caption.style = "Figure Caption"
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER


def replace_references(doc):
    start = find_paragraph(doc, "Références bibliographiques")
    annex = find_paragraph(doc, "Annexes")
    node = start._p.getnext()
    while node is not None and node is not annex._p:
        nxt = node.getnext()
        doc._element.body.remove(node)
        node = nxt

    def builder():
        refs = [
            "[1] Spring, Spring Boot Reference Documentation, https://docs.spring.io/spring-boot/ (consulté le 14/09/2026).",
            "[2] Meta, React Documentation, https://react.dev/ (consultée le 14/09/2026).",
            "[3] Oracle, MySQL Reference Manual, https://dev.mysql.com/doc/ (consulté le 14/09/2026).",
            "[4] Vite, Guide officiel, https://vite.dev/guide/ (consulté le 14/09/2026).",
            "[5] Three.js, Documentation officielle, https://threejs.org/docs/ (consultée le 14/09/2026).",
            "[6] Microsoft, Playwright Documentation, https://playwright.dev/docs/intro (consultée le 14/09/2026).",
            "[7] OWASP Foundation, Cheat Sheet Series, https://cheatsheetseries.owasp.org/ (consultée le 14/09/2026).",
            "[8] W3C, Web Authentication: An API for accessing Public Key Credentials, https://www.w3.org/TR/webauthn-3/ (consulté le 14/09/2026).",
            "[9] Wieland, Products and solutions, https://www.wieland.com/en/products (consulté le 14/09/2026).",
            "[10] KME, Download center, https://www.kme.com/en/downloads (consulté le 14/09/2026).",
            "[11] Aurubis, Products and services, https://www.aurubis.com/en/products (consulté le 14/09/2026).",
            "[12] Nexans, Product catalogue, https://www.nexans.com/en/products.html (consulté le 14/09/2026).",
            "[13] Morocco Copper Foundry, Brochure institutionnelle G-ROD, document fourni avec le projet, 2026.",
            "[14] Projet G-ROD, code source, README, tests et audit technique de la Release Candidate v1.0.0-rc.1, dépôt local, 2026.",
        ]
        for ref in refs:
            p = doc.add_paragraph(ref)
            p.paragraph_format.left_indent = Cm(0.7)
            p.paragraph_format.first_line_indent = Cm(-0.7)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for run in p.runs:
                set_font(run, size=10.5)
        doc.add_page_break()

    append_before(doc, annex, builder)


def configure_document(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(6)
    for style_name, size, color in [
        ("Heading 1", 16, DARK_GREEN),
        ("Heading 2", 14, GREEN),
        ("Heading 3", 12, COPPER),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
    for style_name in ("Figure Caption", "Table Caption"):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(10)
        style.paragraph_format.line_spacing = 1
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style.paragraph_format.keep_with_next = True
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2.5)
        section.header_distance = Cm(1.25)
        section.footer_distance = Cm(1.25)
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def replace_personal_placeholders(doc):
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if "Vous pouvez personnaliser cette dédicace" in text:
            paragraph.text = "[À PERSONNALISER PAR L’ÉTUDIANT]"
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if text.startswith("Je tiens à remercier") and "encadr" in text.lower():
            paragraph.text = ("Je tiens à remercier [TUTEUR ENTREPRISE] pour son accompagnement au sein de Morocco Copper Foundry, "
                              "ainsi que [TUTEUR EMSI] pour son suivi académique. Je remercie également les personnes ayant contribué "
                              "à la compréhension du besoin et à la validation du projet. [À PERSONNALISER PAR L’ÉTUDIANT]")


def patch_cover_and_metadata(path):
    temp = path.with_suffix(".patched.docx")
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    replacements = {
        "Tuteur de l’école : Pr. Prénom et nom": "Tuteur de l’école : [TUTEUR EMSI]",
        "Tuteur de stage : M. ou Mme Prénom et nom": "Tuteur de stage : [TUTEUR ENTREPRISE]",
        "Activité : ---": "Activité : [ACTIVITÉ À VALIDER]",
        "Adresse : ---": "Adresse : [ADRESSE À VALIDER]",
        "Prénom et nom de l’étudiant": "[NOM COMPLET]",
    }
    with zipfile.ZipFile(path, "r") as src, zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "word/document.xml":
                from lxml import etree
                root = etree.fromstring(data)
                for p in root.xpath("//w:txbxContent/w:p", namespaces=ns):
                    nodes = p.xpath(".//w:t", namespaces=ns)
                    combined = "".join(n.text or "" for n in nodes).strip()
                    for old, new in replacements.items():
                        if combined == old and nodes:
                            nodes[0].text = new
                            for node in nodes[1:]:
                                node.text = ""
                            break
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif info.filename == "docProps/core.xml":
                from lxml import etree
                root = etree.fromstring(data)
                cns = {"dc": "http://purl.org/dc/elements/1.1/", "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"}
                creator = root.find("dc:creator", cns)
                if creator is not None:
                    creator.text = "[NOM COMPLET]"
                last = root.find("cp:lastModifiedBy", cns)
                if last is not None:
                    last.text = "[NOM COMPLET]"
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            dst.writestr(info, data)
    temp.replace(path)


def validate(doc, path):
    text = "\n".join(p.text for p in doc.paragraphs)
    table_text = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    combined = text + "\n" + table_text
    banned = ["ZONE RÉSERVÉE", "Illustration manquante", "38/38", "v1.0.0-rc.2"]
    found = [value for value in banned if value in combined]
    if found:
        raise RuntimeError(f"Contenu provisoire ou inexact encore présent : {found}")
    required = ["44/44", "37/38", "Tableau 2.2 — Synthèse du benchmark", "Figure 3.6", "[À PERSONNALISER PAR L’ÉTUDIANT]"]
    missing = [value for value in required if value not in combined]
    if missing:
        raise RuntimeError(f"Éléments requis absents : {missing}")
    if len(doc.sections) != 3:
        raise RuntimeError(f"Trois sections attendues, obtenu : {len(doc.sections)}")
    print(f"Validation structurelle réussie : {len(doc.paragraphs)} paragraphes, {len(doc.tables)} tableaux, {len(doc.inline_shapes)} images.")
    print(path)


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    doc = Document(SOURCE)
    configure_document(doc)
    replace_personal_placeholders(doc)
    replace_summary_sections(doc)
    enrich_introduction(doc)
    enrich_chapter1(doc)
    add_benchmark(doc)
    add_technical_comparisons(doc)
    add_design_analysis(doc)
    fix_known_inaccuracies(doc)
    replace_entity_table(doc)
    enrich_tests_and_limits(doc)
    replace_placeholders(doc)
    replace_references(doc)
    doc.core_properties.title = "Rapport PFE — Plateforme G-ROD"
    doc.core_properties.subject = "Conception et développement d’une plateforme web B2B pour Morocco Copper Foundry"
    doc.core_properties.author = "[NOM COMPLET]"
    doc.core_properties.comments = "Version finale académique — informations personnelles à compléter avant dépôt."
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    patch_cover_and_metadata(OUTPUT)
    final_doc = Document(OUTPUT)
    validate(final_doc, OUTPUT)


if __name__ == "__main__":
    main()
